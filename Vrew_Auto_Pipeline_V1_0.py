#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
금광롱폼 Vrew 자동 파이프라인 v1.0
TTS 음성이 들어간 Vrew 프로젝트(.vrew, v16) 하나를 넣으면 이미지가 연결된 Vrew 파일이 나온다.

  0) login       이멀플(Google Flow) 로그인을 맨 앞에서 확인/완료 — 사람 손이 필요한 일은 여기서 끝나고
                 이후 단계는 무인으로 돈다 (저장된 로그인이 있으면 자동 통과, --no-login-check 로 생략)
  1) extract     Vrew → 클립정보 엑셀              Vrew_Connector_V1_8.extract_clips_from_vrew
  2) storyboard  클립정보 → 스토리보드 엑셀          Claude (스토리보드 프롬프트 마법사 규칙 문서 적용)
  3) characters  등장인물 참조 이미지(캐릭터 시트)    EasyMultyFlow_v5_6 CLI — 분석 JSON 의 외형 묘사로 1인 1장,
                 없는 인물만 생성. 사용자가 characters/<이름>.png 를 넣어 두면 그것을 우선 사용 (--no-characters 로 생략)
  4) images      스토리보드 → Image/ 이미지 생성      EasyMultyFlow_v5_6 CLI (flow.google.com RPC, 기본 — 배치·참조 첨부)
                 / FlowUI_Gen_V1_0 (화면 자동화 폴백) / Multi_Genspark_V3_3 CLI
  5) connect     이미지(·직접 넣은 동영상) → Vrew 클립 연결   Vrew_Connector_V1_8.add_images_to_vrew

사용법
  python Vrew_Auto_Pipeline_V1_0.py init "어디든/프로젝트.vrew"          projects/<오늘날짜>_<이름>/ 폴더를 만들고 파일을 넣는다
  python Vrew_Auto_Pipeline_V1_0.py run "프로젝트.vrew"                 전체 실행
  python Vrew_Auto_Pipeline_V1_0.py run "프로젝트.vrew" --only connect  특정 단계만
  python Vrew_Auto_Pipeline_V1_0.py status "프로젝트.vrew"              진행 상태
  python Vrew_Auto_Pipeline_V1_0.py login                              이멀플(Google Flow) 로그인
  python Vrew_Auto_Pipeline_V1_0.py ui                                 금광 스튜디오 UI (tools/Studio_UI_V1_0.py, 127.0.0.1:7788)
  python Vrew_Auto_Pipeline_V1_0.py progress "projects/<이름>/<이름>.vrew" --watch 30   30초마다 진행 한 줄 (클로드코드가 대화창에 옮겨 적음)
    → 프로젝트별 설정은 projects/<이름>/<이름>_settings.json 에 저장되고 run 이 자동으로 읽는다(CLI 인자가 우선)

같은 명령을 다시 실행하면 끝난 단계는 건너뛴다 (클립정보/스토리보드 엑셀이 있으면 재사용,
이미 만들어진 이미지 번호는 이멀플이 스스로 건너뜀). 처음부터 다시 하려면 --force.

폴더 구성 (이 파일 = 프로젝트 루트)
  tools/        도구 스크립트 13개 (서로 import 하므로 한 폴더에 둔다)   packaging/  PyInstaller spec·ico
  prompts/      스토리보드 프롬프트 마법사 txt                          characters/ 캐릭터 참조 이미지 <영문이름>.png
  projects/     브루 프로젝트별 폴더 <날짜>_<이름>/ (init 이 만든다)
  runtime/      setup_runtime.py 가 설치한 독립 파이썬·라이브러리·Chromium. 있으면 이 스크립트가 자동으로 그 파이썬으로 재실행된다

산출물 (Vrew 파일과 같은 폴더)
  <이름>_클립정보.xlsx        클립번호/시작시간/끝시간/대본 (5~8열 비어 있음)
  <이름>_스토리보드.xlsx      8열 모두 채워진 스토리보드 (이미지 번호·등장인물·프롬프트·켄번즈)
  <이름>_스토리보드_분석.json  Claude 가 확정한 시대배경·캐릭터 묘사·의상 타임라인 (재실행 시 재사용)
  characters/                자동 생성된 캐릭터 시트 <이름>.jpg (공용 characters/ 의 같은 인물 이미지도 여기로 복사)
  CharacterSheets/           캐릭터 시트 원본 출력 (+ <이름>_캐릭터시트.xlsx)
  Image/                     <이미지번호>_<타임스탬프>.jpg  (직접 만든 동영상을 <이미지번호>_*.mp4 로 넣으면 그 번호는 동영상이 연결된다)
  <이름>_이미지연결_<시각>.vrew  최종 결과 (Vrew 에서 열기)
  <이름>_pipeline.json / .log  진행 상태·로그

스토리보드 백엔드 (--storyboard-backend, 기본 auto)
  api         Anthropic API(anthropic 패키지 + ANTHROPIC_API_KEY 또는 ANTHROPIC_AUTH_TOKEN). 키가 있으면 auto 가 이걸 쓴다.
  manual      <이름>_스토리보드_요청.txt 를 만들어 두고 종료(코드 3). 클로드코드 세션이(또는 사람이)
              Claude 에 넣어 이미지 그룹 JSON 배열을 <이름>_스토리보드.json 으로 저장한 뒤 다시 run 하면 이어간다.
              키가 없으면 auto 가 이걸 쓴다 — 파이프라인 기본 동작이다.
  claude CLI(claude -p) 백엔드는 쓰지 않는다 (2026-09-12 사용자 요청으로 제거).
  스토리보드 한 번만 채우면 그 뒤 characters → images → connect 는 무인으로 끝까지 돈다.
  Claude 와의 교환 형식(단계 4)은 JSON — 이미지(그룹)당 {img, from_clip, to_clip, chars, prompt, kenburns} 하나.
  클립번호·시간·대본은 파이프라인이 원본에서 채우므로 Claude 가 되풀이하지 않는다(토큰 절약·옮겨 적기 오류 없음).

동영상은 파이프라인이 만들지 않는다(Grok 자동 생성은 2026-09-11 제거). 직접 만든 동영상을 Image/ 에 <이미지번호>_*.mp4 로
넣으면(스튜디오 UI 스토리보드 카드에서 넣을 수 있다) connect 단계가 그 번호는 이미지 대신 동영상을 연결한다.
"""

__version__ = "1.0"

import os
import sys
import re
import json
import time
import shutil
import argparse
import subprocess
import traceback
from datetime import datetime
from pathlib import Path

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))          # 프로젝트 루트 (이 파일 위치)
TOOLS_DIR = os.path.join(SCRIPT_DIR, "tools")                     # 도구 스크립트 13개
PROMPTS_DIR = os.path.join(SCRIPT_DIR, "prompts")                 # 스토리보드 프롬프트 마법사 txt
if not os.path.isdir(TOOLS_DIR):                                  # 평면 배치(도구가 루트에 있는 경우)도 허용
    TOOLS_DIR = SCRIPT_DIR
if not os.path.isdir(PROMPTS_DIR):
    PROMPTS_DIR = SCRIPT_DIR
if TOOLS_DIR not in sys.path:
    sys.path.insert(0, TOOLS_DIR)
RUNTIME_DIR = os.path.join(SCRIPT_DIR, "runtime")                 # setup_runtime.py 가 만든 독립 파이썬/브라우저
PROJECTS_DIR = os.path.join(SCRIPT_DIR, "projects")


def find_runtime_python():
    for p in (os.path.join(RUNTIME_DIR, "venv", "Scripts", "python.exe"),
              os.path.join(RUNTIME_DIR, "python", "python.exe")):
        if os.path.isfile(p):
            return p
    return ""


def apply_runtime_env():
    """runtime/ms-playwright 가 있으면 이멀플(Playwright)이 그 브라우저를 쓰도록 한다."""
    browsers = os.path.join(RUNTIME_DIR, "ms-playwright")
    if os.path.isdir(browsers) and not os.environ.get("PLAYWRIGHT_BROWSERS_PATH"):
        os.environ["PLAYWRIGHT_BROWSERS_PATH"] = browsers


def reexec_in_runtime():
    """시스템 파이썬으로 실행됐고 runtime/ 파이썬이 있으면 그것으로 다시 실행한다 (라이브러리는 runtime/ 에만 있다)."""
    rp = find_runtime_python()
    if not rp or os.environ.get("VREW_PIPELINE_NO_REEXEC"):
        return
    try:
        if os.path.samefile(rp, sys.executable):
            return
    except OSError:
        pass
    env = dict(os.environ)
    env["VREW_PIPELINE_NO_REEXEC"] = "1"
    env.setdefault("PYTHONIOENCODING", "utf-8")
    env.setdefault("PYTHONUTF8", "1")
    rc = subprocess.call([rp, os.path.abspath(__file__)] + sys.argv[1:], env=env)
    sys.exit(rc)

for _name in ("stdout", "stderr"):
    _stream = getattr(sys, _name, None)
    if _stream is not None and hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

FIXED_COLUMNS = ['클립번호', '시작시간', '끝시간', '대본', '이미지 번호', '등장인물', '프롬프트', '켄번즈']
KENBURNS_VALUES = ("top-to-bottom", "bottom-to-top", "left-to-right", "right-to-left", "zoom-in", "zoom-out")
STAGES = ["login", "extract", "storyboard", "characters", "images", "connect"]
IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.webp'}
VIDEO_EXTS = {'.mp4', '.mov', '.avi', '.mkv', '.webm', '.wmv', '.flv', '.m4v'}
EXIT_WAIT_MANUAL = 3

EMF_SCRIPT = "EasyMultyFlow_v5_6.py"      # 2026-09: flow.google.com batchexecute RPC (로그인 감지·생성·참조 업로드 복구)
FLOWUI_SCRIPT = "FlowUI_Gen_V1_0.py"      # flow.google.com 화면 자동화 — 이멀플 RPC 형식이 또 바뀌었을 때의 폴백(--image-backend flowui)
GENSPARK_SCRIPT = "Multi_Genspark_V3_3.py"
EMF_RATIOS = {"16:9": "16:9 가로 모드", "9:16": "9:16 세로 모드", "1:1": "1:1 정사각"}
IMAGE_RATIO = "16:9"             # 롱폼 전용 — 이미지 비율은 16:9 고정 (캐릭터 시트만 --character-ratio)
DEFAULT_API_MODEL = "claude-opus-5"
DEFAULT_BLOCK_SECONDS = 24.0      # 마법사 예시: 60분 ÷ 150장 ≈ 24초
DEFAULT_CHUNK_SIZE = 120          # Claude 한 번 호출에 넣는 클립 수


class PipelineError(Exception):
    pass


# ══════════════════════════════════════════════════════════════
#  공통 유틸
# ══════════════════════════════════════════════════════════════

def now_str():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def srt_to_seconds(t):
    m = re.match(r"^\s*(\d+):(\d{2}):(\d{2})[,.](\d{1,3})\s*$", str(t or ""))
    if not m:
        m2 = re.match(r"^\s*(\d+):(\d{2}):(\d{2})\s*$", str(t or ""))
        if not m2:
            return 0.0
        h, mi, s = m2.groups()
        return int(h) * 3600 + int(mi) * 60 + int(s)
    h, mi, s, ms = m.groups()
    return int(h) * 3600 + int(mi) * 60 + int(s) + int(ms.ljust(3, "0")) / 1000.0


def fmt_duration(sec):
    sec = int(round(sec))
    return f"{sec // 3600:02d}:{(sec % 3600) // 60:02d}:{sec % 60:02d}"


def find_wizard_file():
    cands = sorted(Path(PROMPTS_DIR).glob("*스토리보드_프롬프트_마법사*.txt"))
    if not cands:
        raise PipelineError(f"스토리보드 프롬프트 마법사 txt 파일을 찾을 수 없습니다: {PROMPTS_DIR}")
    return str(cands[-1])


def scan_media(folder):
    """폴더 안 '<번호>_...' 파일을 번호별로 모은다. {번호: {"image": [경로], "video": [경로]}} (mtime 오름차순)"""
    res = {}
    if not os.path.isdir(folder):
        return res
    for f in os.listdir(folder):
        m = re.match(r"^(\d+)_", f)
        if not m:
            continue
        ext = os.path.splitext(f)[1].lower()
        kind = "image" if ext in IMAGE_EXTS else ("video" if ext in VIDEO_EXTS else None)
        if not kind:
            continue
        res.setdefault(int(m.group(1)), {"image": [], "video": []})[kind].append(os.path.join(folder, f))
    for entry in res.values():
        for k in entry:
            entry[k].sort(key=lambda p: os.path.getmtime(p))
    return res


def parse_number_set(spec):
    """'1,4,7-9' → {1,4,7,8,9}"""
    out = set()
    for part in (spec or "").split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            lo, hi = part.split("-", 1)
            out.update(range(int(lo), int(hi) + 1))
        else:
            out.add(int(part))
    return out


def run_streamed(cmd, cwd=None, env=None, log=print, timeout=None):
    """자식 프로세스를 실행하며 출력을 실시간으로 로그에 흘리고, (종료코드, 전체출력) 을 돌려준다."""
    child_env = dict(os.environ if env is None else env)
    child_env.setdefault("PYTHONIOENCODING", "utf-8")
    child_env.setdefault("PYTHONUTF8", "1")
    log("  $ " + " ".join(f'"{c}"' if " " in str(c) else str(c) for c in cmd))
    proc = subprocess.Popen(cmd, cwd=cwd, env=child_env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            text=True, encoding="utf-8", errors="replace", bufsize=1)
    lines = []
    start = time.time()
    try:
        for line in proc.stdout:
            line = line.rstrip("\r\n")
            lines.append(line)
            log("  │ " + line)
            if timeout and time.time() - start > timeout:
                proc.kill()
                raise PipelineError(f"시간 초과({timeout}초): {cmd[0]}")
    finally:
        proc.stdout.close()
    rc = proc.wait()
    return rc, "\n".join(lines)


def python_has(*modules):
    import importlib.util
    return all(importlib.util.find_spec(m) is not None for m in modules)


# ══════════════════════════════════════════════════════════════
#  엑셀 읽기/쓰기 (8열 스토리보드 규약)
# ══════════════════════════════════════════════════════════════

def read_table_rows(xlsx_path):
    """헤더 기반으로 8열 스토리보드/클립정보 엑셀을 읽는다."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise PipelineError(f"빈 엑셀: {xlsx_path}")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]

    def find(*keys):
        for i, h in enumerate(headers):
            hn = h.replace(" ", "").lower()
            if any(k in hn for k in keys):
                return i
        return None

    ci = {
        "clip_no": find("클립번호", "clip"),
        "start": find("시작", "start"),
        "end": find("끝", "end"),
        "text": find("대본", "script", "자막", "text"),
        "img_num": find("이미지번호", "이미지", "img", "image"),
        "chars": find("등장인물", "character"),
        "prompt": find("프롬프트", "prompt"),
        "kenburns": find("켄번즈", "캔번즈", "켄번스", "kenburns"),
    }
    if ci["clip_no"] is None:
        raise PipelineError(f"'클립번호' 열을 찾을 수 없습니다: {headers}")

    out = []
    for r in rows[1:]:
        def g(k):
            i = ci[k]
            if i is None or i >= len(r) or r[i] is None:
                return ""
            return str(r[i]).strip()
        cn = g("clip_no")
        if not cn:
            continue
        try:
            cn = int(float(cn))
        except ValueError:
            continue
        img_raw = g("img_num")
        img = int(float(img_raw)) if re.match(r"^\d+(\.0+)?$", img_raw) else 0
        out.append({"clip_no": cn, "start": g("start"), "end": g("end"), "text": g("text"),
                    "img_num": img, "chars": g("chars"), "prompt": g("prompt"), "kenburns": g("kenburns")})
    out.sort(key=lambda x: x["clip_no"])
    return out


def write_storyboard_xlsx(path, rows):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "스토리보드"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col, header in enumerate(FIXED_COLUMNS, 1):
        c = ws.cell(row=1, column=col, value=header)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center')
    center = Alignment(horizontal='center')
    for i, r in enumerate(rows):
        row = i + 2
        ws.cell(row=row, column=1, value=r["clip_no"]).alignment = center
        ws.cell(row=row, column=2, value=r["start"])
        ws.cell(row=row, column=3, value=r["end"])
        ws.cell(row=row, column=4, value=r["text"])
        ws.cell(row=row, column=5, value=r["img_num"] if r["img_num"] else None).alignment = center
        ws.cell(row=row, column=6, value=r["chars"] or None)
        ws.cell(row=row, column=7, value=r["prompt"] or None)
        ws.cell(row=row, column=8, value=r["kenburns"] or None)
    for col, w in zip("ABCDEFGH", (10, 18, 13, 50, 12, 25, 80, 15)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def unique_image_numbers(rows):
    return sorted({r["img_num"] for r in rows if r["img_num"]})


# ══════════════════════════════════════════════════════════════
#  스토리보드 검증
# ══════════════════════════════════════════════════════════════

def normalize_chars(raw):
    raw = (raw or "").strip()
    if not raw:
        return ""
    for sep in ("/", "&", ";"):
        raw = raw.replace(sep, ",")
    names = [re.sub(r"\s+", "_", n.strip()) for n in raw.split(",") if n.strip()]
    return ", ".join(dict.fromkeys(names))


def normalize_kenburns(raw):
    v = (raw or "").strip().lower().replace("_", "-").replace(" ", "-")
    if v in KENBURNS_VALUES or v == "none":
        return v
    aliases = {"zoomin": "zoom-in", "zoomout": "zoom-out", "": ""}
    return aliases.get(v.replace("-", ""), "")


def validate_storyboard(rows, clips, start_img=None, strict_contiguous=True):
    """rows(스토리보드) 가 clips(원본 클립) 를 빠짐없이 덮고 규약을 지키는지 검사.
    반환 (정규화된 rows 또는 None, errors, warnings)"""
    errors, warnings = [], []
    by_no = {}
    for r in rows:
        by_no.setdefault(r["clip_no"], r)
    expected = [c["clip_no"] for c in clips]
    missing = [n for n in expected if n not in by_no]
    extra = sorted(set(by_no) - set(expected))
    if missing:
        errors.append(f"누락된 클립 {len(missing)}개: {missing[:15]}{' ...' if len(missing) > 15 else ''}")
    if extra:
        errors.append(f"원본에 없는 클립 {len(extra)}개: {extra[:15]}{' ...' if len(extra) > 15 else ''}")
    if errors:
        return None, errors, warnings

    clip_map = {c["clip_no"]: c for c in clips}
    out = []
    prev_img = None
    for n in expected:
        r = by_no[n]
        img = r["img_num"]
        if not img:
            errors.append(f"클립 {n}: 이미지 번호 없음")
            continue
        if prev_img is None:
            if start_img is not None and img != start_img:
                errors.append(f"첫 이미지 번호는 {start_img}이어야 하는데 {img}")
        elif img < prev_img:
            errors.append(f"클립 {n}: 이미지 번호가 감소 ({prev_img}→{img})")
        elif img > prev_img + 1:
            msg = f"클립 {n}: 이미지 번호 건너뜀 ({prev_img}→{img})"
            (errors if strict_contiguous else warnings).append(msg)
        prev_img = img
        src = clip_map[n]
        out.append({"clip_no": n, "start": src["start"], "end": src["end"], "text": src["text"],
                    "img_num": img, "chars": normalize_chars(r["chars"]),
                    "prompt": (r["prompt"] or "").strip().replace("\n", " "),
                    "kenburns": normalize_kenburns(r["kenburns"])})
        if r["kenburns"] and not out[-1]["kenburns"]:
            warnings.append(f"클립 {n}: 알 수 없는 켄번즈 '{r['kenburns']}' → 빈칸")

    groups = {}
    for r in out:
        groups.setdefault(r["img_num"], []).append(r)
    for img, rs in groups.items():
        first = rs[0]
        if not first["prompt"]:
            donor = next((x for x in rs if x["prompt"]), None)
            if donor:
                first["prompt"] = donor["prompt"]
                first["chars"] = first["chars"] or donor["chars"]
                first["kenburns"] = first["kenburns"] or donor["kenburns"]
            else:
                errors.append(f"이미지 {img}: 프롬프트 없음 (클립 {rs[0]['clip_no']}~{rs[-1]['clip_no']})")
        for x in rs[1:]:
            x["prompt"], x["chars"], x["kenburns"] = "", "", ""
    if errors:
        return None, errors, warnings
    return out, errors, warnings


# ══════════════════════════════════════════════════════════════
#  프로젝트 컨텍스트
# ══════════════════════════════════════════════════════════════

class Project:
    def __init__(self, vrew_path, args):
        self.args = args
        self.vrew = os.path.abspath(vrew_path)
        if not os.path.isfile(self.vrew):
            raise PipelineError(f"Vrew 파일이 없습니다: {self.vrew}")
        self.dir = os.path.dirname(self.vrew)
        self.stem = Path(self.vrew).stem
        j = lambda suffix: os.path.join(self.dir, f"{self.stem}{suffix}")
        self.clips_xlsx = j("_클립정보.xlsx")
        self.storyboard_xlsx = j("_스토리보드.xlsx")
        self.storyboard_csv = j("_스토리보드.csv")          # 옛 수동 형식(| 구분) — 호환용
        self.storyboard_json = j("_스토리보드.json")        # 수동(종료 코드 3) 결과: 이미지 그룹 JSON
        self.request_txt = j("_스토리보드_요청.txt")
        self.settings_json = j("_settings.json")            # 스튜디오 UI 설정(마법사 단계 1~3 + 이미지 옵션) — CLI 인자가 우선
        self.analysis_json = j("_스토리보드_분석.json")
        self.character_xlsx = j("_캐릭터시트.xlsx")
        self.sheet_dir = os.path.join(self.dir, "CharacterSheets")
        self.state_json = j("_pipeline.json")
        self.log_file = j("_pipeline.log")
        self.image_dir = os.path.join(self.dir, "Image")
        char = getattr(args, "char_dir", "") or ""
        if not char:
            local = os.path.join(self.dir, "characters")
            char = local if os.path.isdir(local) else os.path.join(SCRIPT_DIR, "characters")
        self.char_dir = char
        self.state = {"vrew": self.vrew, "stages": {}}
        if os.path.exists(self.state_json):
            try:
                with open(self.state_json, "r", encoding="utf-8") as f:
                    self.state = json.load(f)
            except Exception:
                pass
        self._clips = None

    def log(self, msg):
        line = f"[{now_str()}] {msg}"
        print(line, flush=True)
        try:
            with open(self.log_file, "a", encoding="utf-8") as f:
                f.write(line + "\n")
        except Exception:
            pass

    def mark(self, stage, **info):
        self.state.setdefault("stages", {})[stage] = {"done_at": now_str(), **info}
        with open(self.state_json, "w", encoding="utf-8") as f:
            json.dump(self.state, f, ensure_ascii=False, indent=2)

    def clips(self):
        if self._clips is None:
            self._clips = read_table_rows(self.clips_xlsx)
        return self._clips

    def total_seconds(self):
        clips = self.clips()
        return srt_to_seconds(clips[-1]["end"]) if clips else 0.0


# ══════════════════════════════════════════════════════════════
#  STAGE 0  login — 사람 손이 필요한 로그인을 맨 앞에서 끝낸다 (이후는 무인)
# ══════════════════════════════════════════════════════════════

def stage_login(P, plan):
    """이멀플(Google Flow) 로그인을 파이프라인 시작 시점에 확인/완료한다.
    저장된 로그인이 있으면 창이 잠깐 떴다 닫히고 자동 통과, 없으면 그 창에서 로그인해야 한다."""
    if P.args.dry_run or P.args.no_login_check:
        P.log("[login] 로그인 확인 생략")
        return
    # characters 단계는 flowui 가 아니면(genspark 포함) 항상 이멀플로 캐릭터 시트를 만든다 — stage_characters 와 같은 규칙
    need_emf = (("images" in plan and P.args.image_backend in ("emf", "flowui"))
                or ("characters" in plan and P.args.image_backend != "flowui"))
    if need_emf:
        P.log("[login] 이멀플(Google Flow) 로그인 확인 — 저장된 로그인이 있으면 자동 통과, 없으면 뜨는 브라우저 창에서 로그인하세요 (최대 5분)")
        emf_login(P, "[login]")
        P.log("[login] 이멀플 로그인 확인 완료")
    if not need_emf:
        P.log("[login] 로그인이 필요한 단계 없음")
    P.mark("login", emf=need_emf)


# ══════════════════════════════════════════════════════════════
#  STAGE 1  extract — Vrew → 클립정보 엑셀
# ══════════════════════════════════════════════════════════════

def stage_extract(P):
    from Vrew_Connector_V1_8 import detect_vrew_version, extract_clips_from_vrew
    if os.path.exists(P.clips_xlsx) and not P.args.force:
        clips = P.clips()
        P.log(f"[extract] 기존 클립정보 재사용: {os.path.basename(P.clips_xlsx)} ({len(clips)}클립, {fmt_duration(P.total_seconds())})")
        return
    ver = detect_vrew_version(P.vrew)
    P.log(f"[extract] Vrew 파싱 버전: v{ver}")
    if ver < 16:
        raise PipelineError(f"Vrew v{ver} 파일입니다. v16 이 필요합니다 → Vrew 에서 [파일] > [다른 이름으로 저장] 후 그 파일로 실행하세요.")
    n = extract_clips_from_vrew(P.vrew, P.clips_xlsx)
    P._clips = None
    clips = P.clips()
    total = P.total_seconds()
    P.log(f"[extract] 클립 {n}개 추출 → {os.path.basename(P.clips_xlsx)} (총 길이 {fmt_duration(total)})")
    if n == 0:
        raise PipelineError("클립이 하나도 없습니다. 자막/음성이 들어간 Vrew 파일인지 확인하세요.")
    if total <= 0:
        raise PipelineError("모든 클립 길이가 0초입니다. TTS 음성이 들어간(단어 타이밍이 있는) Vrew 파일이어야 합니다.")
    empty = sum(1 for c in clips if not c["text"])
    if empty:
        P.log(f"[extract] 대본이 빈 클립 {empty}개 (무음/효과 구간) — 앞뒤 이미지 그룹에 포함됩니다")
    P.mark("extract", clips=n, seconds=round(total, 3))


# ══════════════════════════════════════════════════════════════
#  STAGE 2  storyboard — Claude 로 8열 채우기
# ══════════════════════════════════════════════════════════════

def has_api_key():
    return bool(os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN"))


def resolve_backend(P):
    """자동 생성은 Anthropic API 하나뿐이고, 키가 없으면 manual(요청 txt + 종료 코드 3)이다.
    claude CLI(claude -p) 백엔드는 2026-09-12 사용자 요청으로 제거 — 어떤 경우에도 CLI 를 부르지 않는다.
    옛 명령에 남아 있을 수 있는 --storyboard-backend claude-cli 는 auto 로 취급한다(무인 실행이 인자 하나로 죽지 않게)."""
    want = P.args.storyboard_backend
    if want == "claude-cli":
        P.log("[storyboard] --storyboard-backend claude-cli 는 폐지됐습니다 → auto 로 진행")
        want = "auto"
    if want == "api":
        if not python_has("anthropic"):
            raise PipelineError("anthropic 패키지가 없습니다: pip install anthropic")
        if not has_api_key():
            raise PipelineError("ANTHROPIC_API_KEY(또는 ANTHROPIC_AUTH_TOKEN)가 없습니다. "
                                "키 없이 쓰려면 --storyboard-backend manual (기본값 auto 도 같은 동작).")
        return "api"
    if want == "manual":
        return "manual"
    return "api" if (python_has("anthropic") and has_api_key()) else "manual"


def clips_as_pipe_text(clips):
    lines = ["클립번호|시작시간|끝시간|대본"]
    for c in clips:
        text = (c["text"] or "").replace("|", "｜").replace("\n", " ")
        lines.append(f"{c['clip_no']}|{c['start']}|{c['end']}|{text}")
    return "\n".join(lines)


def extract_code_blocks(text):
    blocks = re.findall(r"```[A-Za-z0-9_-]*[ \t]*\r?\n(.*?)```", text, re.S)
    return "\n".join(blocks) if blocks else text


def parse_pipe_rows(text):
    """Claude/사람이 만든 '| 구분 CSV' 텍스트 → row dict 리스트 (검증 전)"""
    rows = []
    for line in extract_code_blocks(text).splitlines():
        line = line.strip().strip("`")
        if not line or "|" not in line:
            continue
        parts = [p.strip() for p in line.split("|")]
        if parts[0].startswith("클립") or not re.match(r"^\d+$", parts[0]):
            continue
        if len(parts) < 5:
            continue
        while len(parts) > 8 and parts[-1] == "":      # 마법사 예시처럼 뒤에 | 가 하나 더 붙은 경우
            parts.pop()
        if len(parts) > 8:                              # 프롬프트 안에 | 가 있으면 프롬프트로 합침
            parts = parts[:6] + ["|".join(parts[6:-1])] + [parts[-1]]
        parts += [""] * (8 - len(parts))
        img = int(parts[4]) if re.match(r"^\d+$", parts[4]) else 0
        rows.append({"clip_no": int(parts[0]), "start": parts[1], "end": parts[2], "text": parts[3],
                     "img_num": img, "chars": parts[5], "prompt": parts[6], "kenburns": parts[7]})
    return rows


def parse_json_object(text):
    t = extract_code_blocks(text).strip()
    s, e = t.find("{"), t.rfind("}")
    if s < 0 or e <= s:
        raise ValueError("JSON 객체를 찾을 수 없음")
    return json.loads(t[s:e + 1])


def parse_json_array(text):
    """Claude/사람이 만든 JSON 배열 텍스트 → list. 코드블록 안이면 그것을, 아니면 첫 '['~마지막 ']' 를 파싱한다.
    {"images": [...]} 처럼 감싼 객체도 받는다. 실패하면 ValueError(원인 포함)."""
    t = extract_code_blocks(text).strip()
    s, e = t.find("["), t.rfind("]")
    os_, oe = t.find("{"), t.rfind("}")
    if s >= 0 and e > s and (os_ < 0 or s < os_):
        cand = t[s:e + 1]
    elif os_ >= 0 and oe > os_:
        cand = t[os_:oe + 1]
    else:
        raise ValueError("JSON 배열을 찾을 수 없음")
    try:
        data = json.loads(cand)
    except json.JSONDecodeError as ex:
        raise ValueError(f"JSON 문법 오류: {ex.msg} (위치 {ex.pos})")
    if isinstance(data, dict):
        for k in ("images", "groups", "storyboard", "rows", "items"):
            if isinstance(data.get(k), list):
                data = data[k]
                break
    if not isinstance(data, list):
        raise ValueError("JSON 최상위가 배열이 아님")
    return data


def rows_from_image_groups(groups, clips, start_img=1):
    """이미지 그룹 JSON [{img, from_clip, to_clip, chars, prompt, kenburns}, …] → 클립별 row (validate_storyboard 전).
    클립번호·시간·대본은 clips(원본)에서 채운다 — Claude 가 되풀이할 필요가 없고 옮겨 적다 틀릴 일도 없다.
    구간(from_clip~to_clip)은 clips 를 빈틈·겹침 없이 순서대로 덮어야 하고, img 는 순서대로 start_img 부터 다시 매긴다.
    반환 (rows, errors, warnings) — errors 가 있으면 rows 는 []."""
    errors, warnings = [], []
    if not isinstance(groups, list) or not groups:
        return [], ["이미지 그룹 배열이 비어 있음"], warnings
    clip_map = {c["clip_no"]: c for c in clips}
    order = [c["clip_no"] for c in clips]

    def as_int(v):
        if v is None or v == "":
            return None
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return None

    parsed = []
    for i, g in enumerate(groups):
        if not isinstance(g, dict):
            errors.append(f"{i + 1}번째 항목이 객체가 아님")
            continue
        a = next((as_int(g.get(k)) for k in ("from_clip", "from", "start_clip", "first_clip", "clip_from")
                  if as_int(g.get(k)) is not None), None)
        b = next((as_int(g.get(k)) for k in ("to_clip", "to", "end_clip", "last_clip", "clip_to")
                  if as_int(g.get(k)) is not None), None)
        if a is None and isinstance(g.get("clips"), list) and g["clips"]:
            a, b = as_int(g["clips"][0]), as_int(g["clips"][-1])
        if a is None and as_int(g.get("clip")) is not None:
            a = as_int(g.get("clip"))
        if b is None:
            b = a
        if a is None:
            errors.append(f"{i + 1}번째 항목: from_clip 없음")
            continue
        chars = g.get("chars", g.get("characters", ""))
        if isinstance(chars, list):
            chars = ", ".join(str(x) for x in chars)
        parsed.append({"img": next((as_int(g.get(k)) for k in ("img", "image", "img_num", "image_num")
                                    if as_int(g.get(k)) is not None), None),
                       "a": a, "b": b, "chars": str(chars or ""), "prompt": str(g.get("prompt") or ""),
                       "kenburns": str(g.get("kenburns") or g.get("ken_burns") or "")})
    if errors:
        return [], errors, warnings
    parsed.sort(key=lambda x: x["a"])
    if parsed[0]["a"] != order[0]:
        errors.append(f"첫 구간이 클립 {order[0]} 이 아니라 {parsed[0]['a']} 에서 시작")
    for i, g in enumerate(parsed):
        if g["b"] < g["a"]:
            errors.append(f"{i + 1}번째 이미지: 구간 역전 ({g['a']}~{g['b']})")
        if i > 0 and g["a"] != parsed[i - 1]["b"] + 1:
            pb = parsed[i - 1]["b"]
            errors.append(f"클립 {pb}→{g['a']}: 구간이 {'겹침' if g['a'] <= pb else '빈틈'} (이전 구간 끝 {pb}, 다음 시작 {g['a']})")
        for n in (g["a"], g["b"]):
            if n not in clip_map:
                errors.append(f"원본에 없는 클립 {n}")
    if parsed[-1]["b"] != order[-1]:
        errors.append(f"마지막 구간이 클립 {order[-1]} 이 아니라 {parsed[-1]['b']} 에서 끝남")
    if errors:
        return [], errors, warnings
    rows = []
    renumbered = 0
    for i, g in enumerate(parsed):
        img = start_img + i
        if g["img"] is not None and g["img"] != img:
            renumbered += 1
        first = True
        for n in order:
            if n < g["a"] or n > g["b"]:
                continue
            src = clip_map[n]
            rows.append({"clip_no": n, "start": src["start"], "end": src["end"], "text": src["text"], "img_num": img,
                         "chars": g["chars"] if first else "", "prompt": g["prompt"] if first else "",
                         "kenburns": g["kenburns"] if first else ""})
            first = False
    if renumbered:
        warnings.append(f"이미지 번호 {renumbered}개를 순서 기준으로 다시 매김 (시작 {start_img})")
    return rows, errors, warnings


def storyboard_json_spec(first_clip, last_clip, start_img, last_rule):
    """Claude 에게 요구하는 스토리보드 JSON 출력 형식 설명(단계 4 프롬프트·수동 요청 txt 공용)."""
    return (
        "- 출력: 아래 형식의 JSON 배열 하나만 코드블록(```json)에 담아 출력하고, 코드블록 밖에는 아무것도 쓰지 않는다. "
        "클립번호·시간·대본은 되풀이하지 않는다(파이프라인이 원본에서 채운다). 이미지(그룹) 하나당 객체 하나:\n"
        '  [{"img": 1, "from_clip": 1, "to_clip": 3, "chars": "Seola, Old_Man", "prompt": "…", "kenburns": "zoom-out"}, …]\n'
        f"  · from_clip~to_clip: 그 이미지가 덮는 연속 클립 구간(양끝 포함). 첫 객체는 클립 {first_clip} 부터, "
        f"마지막 객체는 클립 {last_clip} 까지, 구간이 빈틈·겹침 없이 순서대로 이어져야 한다.\n"
        f"  · img: {start_img} 부터 1씩 증가. chars: 확정된 영문 이름만(공백은 _), 복수는 쉼표, 인물 없는 배경 장면은 \"\".\n"
        f"  · prompt: 영문 한 줄. kenburns: {' / '.join(KENBURNS_VALUES)} 중 하나. {last_rule}\n"
    )


def build_system_prompt(args):
    """마법사 단계 2~3 의 질문(스타일 프롬프트 · 분할 방식/목표 장수)은 run 전에 사용자에게 물어 CLI 인자로 확정된다(스킬이 묻는다).
    단계 1 의 캐릭터 고유 묘사 블록은 항상 사용한다(등장인물 참조 이미지 첨부와 함께 인물 일관성의 축 — 2026-09-11 사용자 결정으로 선택지 제거).
    Claude 에게는 확정값을 전달하고 질문하지 말라고 지시한다."""
    split = getattr(args, "split_mode", "smart")
    return (
        "당신은 유튜브 롱폼 영상용 스토리보드 작가다. 함께 제공되는 '스토리보드 이미지 프롬프트 마법사' 규칙 문서를 그대로 따르되, "
        "지금은 사람과 대화할 수 없는 자동 배치 실행이다.\n"
        "- 마법사 단계 1~3 의 질문은 이미 사용자에게 물어 확정됐다. 질문하지 말고 아래 확정값을 그대로 적용한다: "
        + {"smart": "스마트 장면 전환 분할", "uniform": "균등 분할(목표 블록 길이 단위, 가장 가까운 클립 경계)",
           "front": "초반 집중 분할(도입부는 분당 장수를 높여 촘촘하게, 이후는 블록 길이대로)"}.get(split, "스마트 장면 전환 분할") + ", "
        + "캐릭터 고유 묘사 블록 사용(Yes), "
        "스타일 고정용 프롬프트는 지시된 값만 사용(없으면 스타일 관련 표현 금지).\n"
        "- 요청된 출력 형식 외의 설명·인사·확인 문장은 절대 쓰지 않는다."
    )


class StoryboardGenerator:
    def __init__(self, P, clips, backend):
        self.P = P
        self.clips = clips
        self.backend = backend
        self.args = P.args
        with open(find_wizard_file(), "r", encoding="utf-8") as f:
            self.wizard = f.read()
        self.total_sec = srt_to_seconds(clips[-1]["end"])
        self.block_sec = float(self.args.block_seconds)
        self.split_mode = getattr(self.args, "split_mode", "smart")               # 마법사 단계 3 질문
        self.front_sec = max(0.0, float(getattr(self.args, "front_minutes", 10.0) or 0) * 60)  # 초반 집중 구간 길이 (front:M:A,rest:B 의 M)
        self.front_per_min = max(0.1, float(getattr(self.args, "front_per_min", 3) or 3))     # 초반 구간 분당 장수 (A)
        self.front_rest_per_min = max(0.1, float(getattr(self.args, "front_rest_per_min", 1) or 1))  # 이후 구간 분당 장수 (B)
        # 밀도의 기준은 언제나 하나여야 한다. 스튜디오는 target_images 와 block_seconds/front 파라미터를
        # 늘 함께 저장하므로, 여기서 둘 다 프롬프트에 넣으면 "약 5장" 과 "블록 약 24초"(= 약 15장) 처럼
        # 서로 만족할 수 없는 지시가 되고 Claude 는 블록 길이 쪽을 따라간다
        # (6분14초·목표 5장인데 24초 블록이라 12장이 나온 사례, 2026-09-12). 그래서 한 쪽으로 맞춘다.
        if self.split_mode == "front":
            # 초반 집중 분할은 분당 장수가 유일한 기준 — target_images 는 쓰지 않는다.
            self.target_total = max(1, int(round(self.expected_images(0.0, self.total_sec))))
        elif self.args.target_images:
            # 장수를 직접 정했으면 블록 길이를 거기서 되계산한다.
            self.target_total = int(self.args.target_images)
            if self.total_sec > 0:
                self.block_sec = self.total_sec / self.target_total
        else:
            self.target_total = max(1, int(round(self.expected_images(0.0, self.total_sec))))
        self.style_prompt = (self.args.style_prompt or "").strip()
        self.system_prompt = build_system_prompt(self.args)
        self.model = self.args.model or DEFAULT_API_MODEL

    def wizard_override_text(self):
        """마법사 규칙 문서는 단계 3 의 기본값을 '150개 · 목표 블록 약 24초 · 12~36초' 라고 못 박아 두었다.
        시스템 프롬프트가 그 문서를 '그대로 따르라' 고 하므로 구간 지시 한 줄("약 5개")보다 문서 쪽이 앞선다 —
        목표 5장인데 31초/장(12장)이 나온 것이 정확히 문서의 12~36초 범위였다 (2026-09-12).
        그래서 문서 바로 뒤에 확정값으로 덮어쓰는 블록을 붙인다."""
        lines = ["[이 실행의 확정값 — 위 규칙 문서 단계 3 의 기본값을 덮어쓴다]",
                 "- 규칙 문서의 '기본값 150개', '목표 블록 길이 약 24초', '12초~36초' 는 그 문서가 든 예시다. 이 실행에는 적용하지 않는다.",
                 f"- 이 영상: 총 길이 {fmt_duration(self.total_sec)}, 목표 이미지 {self.target_total}개."]
        if self.split_mode == "front":
            lines.append(f"- 초반 {self.front_sec / 60:.0f}분은 분당 {self.front_per_min:g}장, 그 뒤는 분당 {self.front_rest_per_min:g}장.")
        else:
            lines.append(f"- 목표 블록 길이 = {fmt_duration(self.total_sec)} ÷ {self.target_total}개 ≈ {self.block_sec:.0f}초 "
                         f"(허용 {self.block_sec * 0.5:.0f}~{self.block_sec * 1.5:.0f}초). 장면 전환은 이 범위 안에서 고르고, "
                         f"더 잘게 쪼개 장수를 늘리지 않는다.")
        lines.append(f"- 최종 출력은 이미지 {self.target_total}개 안팎이어야 한다. 이 숫자가 규칙 문서의 어떤 기본값보다 우선한다.")
        return chr(10).join(lines) + chr(10)

    def expected_images(self, t0, t1):
        """구간 [t0, t1) 초에 기대되는 이미지 수(밀도 적분). smart/uniform = 블록 길이당 1장,
        front = 초반 front_sec 까지는 분당 front_per_min 장, 그 뒤는 분당 front_rest_per_min 장 (라움튜브 front:M:A,rest:B)."""
        t0, t1 = max(0.0, t0), max(0.0, t1)
        if t1 <= t0:
            return 0.0
        if self.split_mode != "front":
            return (t1 - t0) / self.block_sec
        front = max(0.0, min(t1, self.front_sec) - t0)
        rest = (t1 - t0) - front
        return front / 60.0 * self.front_per_min + rest / 60.0 * self.front_rest_per_min

    # ── 백엔드 호출 ──
    def complete(self, user_prompt, label):
        self.P.log(f"[storyboard] Claude 호출 ({self.backend}, {self.model}): {label}")
        t0 = time.time()
        text = self._complete_api(user_prompt)
        self.P.log(f"[storyboard]   응답 {len(text)}자 ({time.time() - t0:.0f}초)")
        return text

    def _complete_api(self, user_prompt):
        import anthropic
        client = anthropic.Anthropic()
        system = [{"type": "text", "text": self.system_prompt},
                  {"type": "text", "text": "[규칙 문서: 스토리보드 이미지 프롬프트 마법사]\n" + self.wizard,
                   "cache_control": {"type": "ephemeral"}},
                  {"type": "text", "text": self.wizard_override_text()}]   # 캐시 밖 — 실행마다 달라진다
        kwargs = dict(model=self.model, max_tokens=64000, system=system,
                      messages=[{"role": "user", "content": user_prompt}],
                      output_config={"effort": "high"})
        try:
            # 안전 분류기 거부 시 서버가 대체 모델로 같은 요청을 이어서 처리한다(server-side fallback).
            with client.beta.messages.stream(betas=["server-side-fallback-2026-06-01"],
                                             fallbacks=[{"model": "claude-opus-4-8"}], **kwargs) as s:
                msg = s.get_final_message()
        except anthropic.BadRequestError as e:
            low = str(e).lower()
            if "fallback" not in low and "beta" not in low:
                raise
            self.P.log("[storyboard]   fallback 파라미터 미지원 → 기본 호출로 재시도")
            with client.messages.stream(**kwargs) as s:
                msg = s.get_final_message()
        if msg.stop_reason == "refusal":
            raise PipelineError("Claude 가 요청을 거부했습니다 (stop_reason=refusal). 대본 내용을 확인하세요.")
        text = "".join(b.text for b in msg.content if b.type == "text")
        if msg.stop_reason == "max_tokens":
            self.P.log("[storyboard]   경고: 출력이 max_tokens 에서 잘렸습니다 — 검증에서 걸리면 청크를 줄이세요(--chunk-size)")
        return text

    # ── 단계 1: 분석 ──
    def analyze(self):
        if os.path.exists(self.P.analysis_json) and not self.args.force:
            with open(self.P.analysis_json, "r", encoding="utf-8") as f:
                data = json.load(f)
            self.P.log(f"[storyboard] 기존 분석 재사용: {os.path.basename(self.P.analysis_json)} "
                       f"(캐릭터 {len(data.get('characters', []))}명)")
            return data
        last = self.clips[-1]["clip_no"]
        schema = (
            '{\n'
            '  "era_setting": "시대·지리적 배경 (한국어 1~2문장)",\n'
            '  "summary": "줄거리 요약 (한국어 3~6문장, 장면 전환과 사건 순서 위주)",\n'
            '  "characters": [\n'
            '    {\n'
            '      "name": "영문 이름 (공백은 _ 로. 예: Seola, Old_Man, Police_Officer)",\n'
            '      "name_ko": "한글 이름/호칭",\n'
            '      "description": "캐릭터 고유 묘사 블록 (영문, 의상 제외: 나이대·성별·체형·얼굴·머리·인상). 모든 프롬프트 앞에 그대로 반복될 문장",\n'
            '      "costume_timeline": [\n'
            '        {"from_clip": 1, "to_clip": 40, "situation": "상황 (한국어)", "costume": "wearing ... (영문)"}\n'
            '      ]\n'
            '    }\n'
            '  ],\n'
            '  "locations": [ {"name": "장소 (한국어)", "description": "구체적 배경 묘사 (영문, 여백 방지 원칙 적용)"} ],\n'
            f'  "style_prompt": {json.dumps(self.style_prompt, ensure_ascii=False)},\n'
            f'  "target_image_count": {self.target_total}\n'
            '}'
        )
        prompt = (
            "[작업: 단계 1 — 대본 분석 및 비주얼 가이드 확정]\n"
            "아래 클립 데이터(전체 대본)를 분석해 다음 형식의 JSON 객체 하나만 출력한다. 코드블록·설명 없이 순수 JSON 만 출력.\n"
            f"{schema}\n"
            "- characters 는 대본에 실제로 등장해 그림에 나올 인물만, 중요한 순서로. 이름은 이후 모든 구간에서 동일하게 쓴다.\n"
            f"- costume_timeline 은 각 인물별로 클립 1~{last} 전체를 빠짐없이 덮도록 구간을 나눈다(의상 변화 단서: 직업·장소·시간·계절·사건).\n"
            f"- 총 영상 길이 {fmt_duration(self.total_sec)}, 목표 이미지 수 {self.target_total}개"
            + ("" if self.split_mode == "front" else f" (목표 블록 길이 약 {self.block_sec:.0f}초)") + ".\n"
            + {"smart": "- 분할 방식(확정): 스마트 장면 전환 분할 — 목표 블록 길이 근처의 장면 전환 지점에서 나눈다(±50%).\n",
               "uniform": "- 분할 방식(확정): 균등 분할 — 목표 블록 길이마다 가장 가까운 클립 경계에서 나눈다(장면 전환보다 간격 우선).\n",
               "front": f"- 분할 방식(확정): 초반 집중 — 영상 시작 {self.front_sec / 60:.1f}분까지는 분당 {self.front_per_min:g}장으로 촘촘하게, "
                        f"그 뒤는 분당 {self.front_rest_per_min:g}장(약 {60 / self.front_rest_per_min:.0f}초당 1장). 장면 전환 지점을 우선하되 구간별 밀도를 지킨다.\n"}[self.split_mode if self.split_mode in ("smart", "uniform", "front") else "smart"]
            + "- 캐릭터 고유 묘사 블록(확정): 사용 — description 을 모든 프롬프트 앞에 그대로 반복한다.\n"
            + ("- 스타일 고정용 프롬프트가 지정되어 있다: 모든 프롬프트 앞에 그대로 넣는다.\n" if self.style_prompt
               else "- 스타일 고정용 프롬프트 없음: 스타일 관련 표현(Cinematic, film still, realistic lighting 등)을 쓰지 않는다.\n")
            + "\n[클립 데이터]\n" + clips_as_pipe_text(self.clips)
        )
        data = None
        errs = []
        for attempt in range(1, 3):
            text = self.complete(prompt if attempt == 1 else prompt + f"\n\n[이전 출력 문제 — 반드시 수정] {errs[-1]}", "단계 1 분석")
            try:
                data = parse_json_object(text)
                if not isinstance(data.get("characters"), list):
                    raise ValueError("characters 배열 없음")
                break
            except Exception as e:
                errs.append(f"JSON 파싱 실패: {e}")
                self.P.log(f"[storyboard]   {errs[-1]} — 재시도")
                data = None
        if data is None:
            raise PipelineError("단계 1 분석 JSON 을 얻지 못했습니다: " + "; ".join(errs))
        for ch in data["characters"]:
            ch["name"] = re.sub(r"\s+", "_", str(ch.get("name", "")).strip())
        data["style_prompt"] = self.style_prompt
        data["target_image_count"] = self.target_total
        data["character_block"] = True
        data["split_mode"] = self.split_mode
        data["block_seconds"] = self.block_sec
        if self.split_mode == "front":
            data["front_minutes"] = self.front_sec / 60.0
            data["front_per_min"] = self.front_per_min
            data["front_rest_per_min"] = self.front_rest_per_min
        with open(self.P.analysis_json, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        names = ", ".join(c["name"] for c in data["characters"]) or "(없음)"
        self.P.log(f"[storyboard] 분석 완료 → {os.path.basename(self.P.analysis_json)} | 캐릭터: {names}")
        return data

    # ── 단계 4: 구간별 이미지 그룹 JSON ──
    def plan_chunks(self):
        size = max(10, int(self.args.chunk_size))
        chunks = [self.clips[i:i + size] for i in range(0, len(self.clips), size)]
        if len(chunks) > 1 and len(chunks[-1]) < size * 0.3:
            chunks[-2].extend(chunks.pop())
        return chunks

    def split_instruction(self, chunk, target_imgs):
        """구간 프롬프트의 분할 지시문(분할 방식별)."""
        if self.split_mode == "uniform":
            return (f"균등 분할: 약 {self.block_sec:.0f}초마다 가장 가까운 클립 경계에서 나눈다(편차 ±20%, 문장이 잘리지 않게). "
                    f"이 구간은 약 {target_imgs}개가 된다. 장면 전환보다 간격을 우선한다.\n")
        if self.split_mode == "front":
            t0, t1 = srt_to_seconds(chunk[0]["start"]), srt_to_seconds(chunk[-1]["end"])
            front = max(0.0, min(t1, self.front_sec) - t0)
            rest = max(0.0, (t1 - t0) - front)
            n_front = int(round(front / 60.0 * self.front_per_min))
            n_rest = max(0, target_imgs - n_front) if front > 0 else target_imgs
            parts = []
            if front > 0:
                parts.append(f"영상 시작 {self.front_sec / 60:.1f}분까지(이 구간의 앞 {front:.0f}초)는 도입부라 분당 {self.front_per_min:g}장으로 촘촘하게 약 {max(1, n_front)}개")
            if rest > 0:
                parts.append(f"그 뒤 {rest:.0f}초는 분당 {self.front_rest_per_min:g}장(약 {60 / self.front_rest_per_min:.0f}초당 1장)으로 약 {max(1, n_rest) if front > 0 else target_imgs}개")
            return ("초반 집중 분할: " + ", ".join(parts) + f" — 합계 약 {target_imgs}개(±30%). 밀도를 지키되 장면 전환 지점을 우선한다.\n")
        return (f"이 구간의 목표 이미지 수는 약 {target_imgs}개(±30%), 목표 블록 길이 약 {self.block_sec:.0f}초(±50%) — 장면 전환 지점을 우선해 분할한다.\n")

    def generate_chunk(self, chunk, analysis, start_img, target_imgs, prev_tail, is_last):
        a, b = chunk[0]["clip_no"], chunk[-1]["clip_no"]
        n = len(chunk)
        chunk_sec = srt_to_seconds(chunk[-1]["end"]) - srt_to_seconds(chunk[0]["start"])
        style_line = (f'스타일 고정용 프롬프트 (모든 프롬프트 맨 앞에 그대로 삽입): "{self.style_prompt}"' if self.style_prompt
                      else "스타일 고정용 프롬프트: 없음 → 스타일 관련 표현 금지")
        last_rule = f"영상 전체의 마지막 클립(클립 {b})의 켄번즈는 반드시 none." if is_last else "이 구간은 영상 중간이므로 none 을 쓰지 않는다."
        prev_block = ""
        if prev_tail:
            prev_block = ("\n[직전 구간의 마지막 이미지 — 연속성 참고용, 다시 출력하지 말 것]\n"
                          f"이미지 {prev_tail['img_num']} (클립 ~{prev_tail['clip_no']}): 등장인물={prev_tail['chars'] or '(없음)'}; "
                          f"프롬프트={prev_tail['prompt']}\n")
        prompt = (
            "[확정된 설정 — 단계 1 결과 (그대로 적용)]\n"
            + json.dumps(analysis, ensure_ascii=False, indent=1)
            + f"\n{style_line}\n\n"
            f"[작업: 단계 4 — 클립 {a}~{b} ({n}개, 약 {chunk_sec:.0f}초) 스토리보드 작성]\n"
            f"- 아래 클립 {a}~{b} 각각에 대해 이미지 번호·등장인물·프롬프트·켄번즈를 채운다.\n"
            f"- 이미지 번호는 {start_img}부터 시작해 1씩 증가한다 (건너뛰기·감소 금지). "
            + self.split_instruction(chunk, target_imgs)
            + f"- 모든 클립({a}~{b}, {n}개)이 정확히 하나의 이미지 구간에 속해야 한다. 같은 이미지를 공유하는 연속 클립은 하나의 객체(from_clip~to_clip)로 묶는다.\n"
            + "- 프롬프트: 규칙 문서 4~6절대로. [캐릭터 고유 묘사 블록 전체 반복] + [해당 구간 의상 (costume_timeline 준수)] + [동작/표정] + "
              "[구체적 배경/장면 (locations 활용, 여백 방지 원칙)], 끝에 no text, no watermark. 영문 한 줄.\n"
            +
            "- 등장인물: 확정된 영문 이름만 사용, 공백은 _, 복수는 쉼표. 인물 없는 배경 장면은 빈칸.\n"
            + storyboard_json_spec(a, b, start_img, last_rule)
            + prev_block
            + "\n[클립 데이터]\n" + clips_as_pipe_text(chunk)
        )
        last_errors = []
        for attempt in range(1, 4):
            p = prompt
            if last_errors:
                p += "\n\n[이전 출력의 문제점 — 반드시 고쳐서 전체를 다시 출력] " + "; ".join(last_errors[:8])
            text = self.complete(p, f"클립 {a}~{b} (시도 {attempt})")
            try:
                rows, gerrors, gwarnings = rows_from_image_groups(parse_json_array(text), chunk, start_img)
            except ValueError as e:
                rows, gerrors, gwarnings = [], [f"JSON 파싱 실패: {e}"], []
            for w in gwarnings[:5]:
                self.P.log(f"[storyboard]   경고: {w}")
            if gerrors:
                last_errors = gerrors
                self.P.log(f"[storyboard]   구간 검증 실패 ({len(gerrors)}건): " + "; ".join(gerrors[:4]))
                continue
            ok, errors, warnings = validate_storyboard(rows, chunk, start_img=start_img, strict_contiguous=True)
            for w in warnings:
                self.P.log(f"[storyboard]   경고: {w}")
            if ok:
                imgs = unique_image_numbers(ok)
                tol = max(1, int(round(target_imgs * 0.3)))       # 프롬프트가 약속한 ±30% (최소 ±1장)
                off = abs(len(imgs) - target_imgs) > tol
                if off and attempt < 3:
                    self.P.log(f"[storyboard]   장수 불일치: {len(imgs)}장 (목표 {target_imgs}±{tol}장) → 재시도")
                    last_errors = [f"이미지를 {len(imgs)}장 냈는데 목표는 {target_imgs}장(±{tol})이다. "
                                   f"구간을 더 {'넓게' if len(imgs) > target_imgs else '잘게'} 잡아 {target_imgs}장에 맞춰 다시 출력"]
                    continue
                if off:
                    self.P.log(f"[storyboard]   경고: 이미지 {len(imgs)}장 — 목표 {target_imgs}±{tol}장에 못 맞춤 (3회 시도, 그대로 진행)")
                self.P.log(f"[storyboard]   클립 {a}~{b} 완료: 이미지 {imgs[0]}~{imgs[-1]} ({len(imgs)}장)")
                return ok
            last_errors = errors
            self.P.log(f"[storyboard]   검증 실패 ({len(errors)}건): " + "; ".join(errors[:4]))
        raise PipelineError(f"클립 {a}~{b} 스토리보드를 3회 시도해도 검증을 통과하지 못했습니다: " + "; ".join(last_errors[:5]))

    def generate(self):
        analysis = self.analyze()
        chunks = self.plan_chunks()
        self.P.log(f"[storyboard] 총 {len(self.clips)}클립 / {fmt_duration(self.total_sec)} → 목표 이미지 {self.target_total}장, "
                   f"{len(chunks)}개 구간으로 생성")
        rows_all = []
        next_img = 1
        prev_tail = None
        for i, chunk in enumerate(chunks):
            t0, t1 = srt_to_seconds(chunk[0]["start"]), srt_to_seconds(chunk[-1]["end"])
            whole = max(self.expected_images(0.0, self.total_sec), 1e-6)
            target = max(1, int(round(self.target_total * self.expected_images(t0, t1) / whole)))
            rows = self.generate_chunk(chunk, analysis, next_img, target, prev_tail, is_last=(i == len(chunks) - 1))
            rows_all.extend(rows)
            groups = {}
            for r in rows:
                groups.setdefault(r["img_num"], []).append(r)
            last_img = max(groups)
            first = groups[last_img][0]
            prev_tail = {"img_num": last_img, "clip_no": rows[-1]["clip_no"], "chars": first["chars"], "prompt": first["prompt"]}
            next_img = last_img + 1
        # 마지막 클립 규칙: 마지막 이미지 그룹의 켄번즈는 none
        last_group_first = next(r for r in rows_all if r["img_num"] == rows_all[-1]["img_num"])
        last_group_first["kenburns"] = "none"
        return rows_all


def write_manual_request(P, clips):
    with open(find_wizard_file(), "r", encoding="utf-8") as f:
        wizard = f.read()
    total = srt_to_seconds(clips[-1]["end"])
    a = P.args
    mode = getattr(a, "split_mode", "smart")
    fm = float(getattr(a, "front_minutes", 10) or 10)
    fa = float(getattr(a, "front_per_min", 3) or 3)
    fb = float(getattr(a, "front_rest_per_min", 1) or 1)
    if mode == "front":
        fsec = min(total, fm * 60)
        auto = fsec / 60 * fa + (total - fsec) / 60 * fb
    else:
        auto = total / float(a.block_seconds)
    target = P.args.target_images or max(1, int(round(auto)))
    split_txt = {"uniform": f"균등 분할 (블록 {float(a.block_seconds):.0f}초 간격)",
                 "front": f"초반 집중 (처음 {fm:g}분은 분당 {fa:g}장, 이후 분당 {fb:g}장 — front:{fm:g}:{fa:g},rest:{fb:g})"}.get(mode, "스마트 장면 전환 분할")
    block_txt = "사용 (항상)"
    style_txt = (a.style_prompt or "").strip() or "없음 (스타일 관련 표현 금지)"
    body = (
        "=== 금광롱폼 파이프라인: 스토리보드 생성 요청 ===\n"
        f"프로젝트: {os.path.basename(P.vrew)}\n"
        f"클립 {len(clips)}개, 총 길이 {fmt_duration(total)}, 목표 이미지 수 약 {target}장\n"
        f"확정 설정 (마법사 단계 1~3 — 사용자가 이미 고른 값, 다시 묻지 않는다):\n"
        f"  · 이미지 분할: {split_txt}, 목표 블록 길이 약 {float(a.block_seconds):.0f}초\n"
        f"  · 캐릭터 고유 묘사 블록: {block_txt}\n"
        f"  · 스타일 고정용 프롬프트: {style_txt}\n\n"
        "사용 방법\n"
        " 1. 아래 '규칙 문서'와 '클립 데이터'를 Claude(claude.ai 또는 클로드코드)에 그대로 붙여 넣는다.\n"
        " 2. 단계별 질문은 위 '확정 설정' 대로 진행하고(질문 불필요), 결과는 아래 'JSON 출력 형식' 으로 받는다\n"
        "    (마법사의 엑셀/CSV 출력 형식 질문은 생략 — 엑셀은 파이프라인이 만든다). 길면 나눠 받은 배열을 하나로 합친다.\n"
        f" 3. 결과 JSON(UTF-8)을 다음 이름으로 저장한다:  {os.path.basename(P.storyboard_json)}\n"
        f"    (엑셀로 받았다면 {os.path.basename(P.storyboard_xlsx)} 로 저장)\n"
        f" 4. 다시 실행:  python {os.path.basename(__file__)} run \"{P.vrew}\"\n\n"
        "─── JSON 출력 형식 ───\n"
        + storyboard_json_spec(clips[0]["clip_no"], clips[-1]["clip_no"], 1,
                               f"영상 전체의 마지막 이미지(클립 {clips[-1]['clip_no']} 포함 구간)의 kenburns 는 반드시 none.")
        + "\n─── 규칙 문서 ───\n" + wizard + "\n\n"
        "─── 클립 데이터 (클립번호|시작시간|끝시간|대본) ───\n" + clips_as_pipe_text(clips) + "\n"
    )
    with open(P.request_txt, "w", encoding="utf-8") as f:
        f.write(body)


def stage_storyboard(P):
    clips = P.clips()
    if os.path.exists(P.storyboard_xlsx) and not P.args.force:
        rows = read_table_rows(P.storyboard_xlsx)
        ok, errors, warnings = validate_storyboard(rows, clips, strict_contiguous=False)
        if ok:
            imgs = unique_image_numbers(ok)
            P.log(f"[storyboard] 기존 스토리보드 재사용: {os.path.basename(P.storyboard_xlsx)} (이미지 {len(imgs)}장, {imgs[0]}~{imgs[-1]})")
            for w in warnings[:10]:
                P.log(f"[storyboard]   경고: {w}")
            return ok
        P.log(f"[storyboard] 기존 스토리보드가 규약에 맞지 않습니다: " + "; ".join(errors[:5]))
        bak = P.storyboard_xlsx + ".bak"
        shutil.move(P.storyboard_xlsx, bak)
        P.log(f"[storyboard]   → {os.path.basename(bak)} 으로 옮기고 새로 생성합니다")

    if os.path.exists(P.storyboard_json):
        with open(P.storyboard_json, "r", encoding="utf-8-sig") as f:
            text = f.read()
        try:
            rows, gerrors, gwarnings = rows_from_image_groups(parse_json_array(text), clips, start_img=1)
        except ValueError as e:
            raise PipelineError(f"{os.path.basename(P.storyboard_json)} 파싱 실패: {e}")
        if gerrors:
            raise PipelineError(f"{os.path.basename(P.storyboard_json)} 구간 검증 실패: " + "; ".join(gerrors[:6]))
        for w in gwarnings[:5]:
            P.log(f"[storyboard]   경고: {w}")
        ok, errors, warnings = validate_storyboard(rows, clips, start_img=1, strict_contiguous=True)
        if not ok:
            raise PipelineError(f"{os.path.basename(P.storyboard_json)} 검증 실패: " + "; ".join(errors[:6]))
        for w in warnings[:10]:
            P.log(f"[storyboard]   경고: {w}")
        last_first = next(r for r in ok if r["img_num"] == ok[-1]["img_num"])
        last_first["kenburns"] = "none"
        write_storyboard_xlsx(P.storyboard_xlsx, ok)
        imgs = unique_image_numbers(ok)
        P.log(f"[storyboard] JSON → 엑셀 변환 완료: {os.path.basename(P.storyboard_xlsx)} (이미지 {len(imgs)}장)")
        P.mark("storyboard", source="json", images=len(imgs))
        return ok

    if os.path.exists(P.storyboard_csv):          # 옛 수동 형식(| 구분 CSV) 호환
        with open(P.storyboard_csv, "r", encoding="utf-8-sig") as f:
            text = f.read()
        rows = parse_pipe_rows(text)
        ok, errors, warnings = validate_storyboard(rows, clips, strict_contiguous=False)
        if not ok:
            raise PipelineError(f"{os.path.basename(P.storyboard_csv)} 검증 실패: " + "; ".join(errors[:6]))
        for w in warnings[:10]:
            P.log(f"[storyboard]   경고: {w}")
        write_storyboard_xlsx(P.storyboard_xlsx, ok)
        imgs = unique_image_numbers(ok)
        P.log(f"[storyboard] CSV → 엑셀 변환 완료: {os.path.basename(P.storyboard_xlsx)} (이미지 {len(imgs)}장)")
        P.mark("storyboard", source="csv", images=len(imgs))
        return ok

    backend = resolve_backend(P)
    if P.args.dry_run and backend != "manual":
        raise PipelineError("[dry-run] 스토리보드가 없어 Claude 호출이 필요합니다. dry-run 을 빼고 실행하세요.")
    if backend == "manual":
        write_manual_request(P, clips)
        P.log("[storyboard] 스토리보드를 클로드코드 세션이 작성합니다 (API 키 없음 → manual) → 요청 파일 작성")
        P.log(f"[storyboard]   {P.request_txt}")
        P.log(f"[storyboard]   Claude 결과(이미지 그룹 JSON 배열)를 {os.path.basename(P.storyboard_json)} 로 저장한 뒤 같은 명령을 다시 실행하세요.")
        P.mark("storyboard", waiting="manual")
        sys.exit(EXIT_WAIT_MANUAL)

    gen = StoryboardGenerator(P, clips, backend)
    rows = gen.generate()
    ok, errors, warnings = validate_storyboard(rows, clips, start_img=1, strict_contiguous=True)
    if not ok:
        raise PipelineError("최종 스토리보드 검증 실패: " + "; ".join(errors[:6]))
    write_storyboard_xlsx(P.storyboard_xlsx, ok)
    try:
        from StoryBoard_Checker_V2_0 import parse_excel as sb_parse_excel, compare_clips as sb_compare
        results = sb_compare(sb_parse_excel(P.clips_xlsx), sb_parse_excel(P.storyboard_xlsx))
        bad = [r for r in results if not r["ok"]]
        P.log(f"[storyboard] 스토리보드 체커: {len(results)}클립 중 불일치 {len(bad)}건")
    except Exception as e:
        P.log(f"[storyboard] 스토리보드 체커 생략: {e}")
    imgs = unique_image_numbers(ok)
    chars = sorted({n.strip() for r in ok for n in (r["chars"] or "").split(",") if n.strip()})
    P.log(f"[storyboard] 완료 → {os.path.basename(P.storyboard_xlsx)} | 이미지 {len(imgs)}장 | 등장인물: {', '.join(chars) or '(없음)'}")
    if chars:
        have = {os.path.splitext(f)[0].lower() for f in os.listdir(P.char_dir)} if os.path.isdir(P.char_dir) else set()
        missing = [c for c in chars if c.lower() not in have]
        if missing:
            P.log(f"[storyboard]   참고: {P.char_dir} 에 참조 이미지가 없는 인물: {', '.join(missing)} "
                  f"(<이름>.png 를 넣으면 이멀플이 캐릭터 에셋으로 자동 첨부)")
    P.mark("storyboard", source=backend, images=len(imgs), characters=chars,
           settings={"style_prompt": (P.args.style_prompt or "").strip(), "target_images": P.args.target_images or None,
                     "block_seconds": float(P.args.block_seconds), "split_mode": getattr(P.args, "split_mode", "smart"),
                     "character_block": True})
    return ok


# ══════════════════════════════════════════════════════════════
#  STAGE 3  images — 이멀플 / Genspark CLI
# ══════════════════════════════════════════════════════════════

def find_tool_runner(script_name):
    """<도구>.exe 가 있으면 그것, 아니면 [현재 파이썬, <도구>.py]"""
    stem = os.path.splitext(script_name)[0]
    for cand in (os.path.join(TOOLS_DIR, stem + ".exe"),
                 os.path.join(SCRIPT_DIR, "dist", stem + ".exe"),
                 os.path.join(SCRIPT_DIR, "packaging", "dist", stem + ".exe")):
        if os.path.isfile(cand):
            return [cand]
    py = os.path.join(TOOLS_DIR, script_name)
    if not os.path.isfile(py):
        raise PipelineError(f"{script_name} 이 도구 폴더에 없습니다: {TOOLS_DIR}")
    return [sys.executable, py]


def flow_tool_runner(script_name):
    """Flow 도구(FlowUI_Gen / 이멀플) 실행 명령 — .py 로 돌릴 때는 필요한 패키지를 먼저 확인해 setup_runtime 안내를 낸다
    (login 단계가 맨 먼저 부르므로 여기서 안 잡으면 ImportError 가 '로그인 실패' 로 잘못 보인다)."""
    runner = find_tool_runner(script_name)
    need = ("playwright", "openpyxl") if script_name == FLOWUI_SCRIPT else ("PySide6", "playwright", "openpyxl")
    if runner[0] == sys.executable and not python_has(*need):
        raise PipelineError(f"{script_name} 실행에 {'/'.join(need)} 가 필요합니다: python setup_runtime.py")
    return runner


def missing_image_numbers(P, rows):
    expected = unique_image_numbers(rows)
    have = scan_media(P.image_dir)
    return [n for n in expected if n not in have or not (have[n]["image"] or have[n]["video"])]


EMF_LOGIN_NEEDED = ("저장된 로그인이 없습니다", "세션 준비 실패", "먼저 로그인", "로그인 안 됨", "로그인 실패")


def emf_command(P, excel, image_dir, *, only=None, ratio=None, dry_run=False):
    """이미지 생성 CLI 명령. image_backend 가 emf 면 이멀플(flow.google.com RPC, 기본), flowui 면 FlowUI_Gen(화면 자동화).
    --headless 는 둘 다 지원. (이멀플 v5.6 은 헤드리스 UA 를 교정해 Flow 의 UNUSUAL_ACTIVITY 거부를 피하고,
    그래도 그 403 이 나오면 스스로 창 모드로 전환해 이어간다.)"""
    a = P.args
    if a.image_backend == "flowui":
        runner = find_tool_runner(FLOWUI_SCRIPT)
        if runner[0] == sys.executable and not python_has("playwright", "openpyxl"):
            raise PipelineError("FlowUI 실행에 playwright 가 필요합니다: python setup_runtime.py")
        cmd = runner + ["--generate", excel, "--image-dir", image_dir, "--model", a.emf_model,
                        "--ratio", (ratio or EMF_RATIOS[IMAGE_RATIO]).split()[0],
                        "--project-store", os.path.join(P.dir, f"{P.stem}_flow_project.json")]
    else:
        runner = find_tool_runner(EMF_SCRIPT)
        if runner[0] == sys.executable and not python_has("PySide6", "playwright", "openpyxl"):
            raise PipelineError("이멀플 실행에 PySide6/playwright 가 필요합니다: python setup_runtime.py")
        cmd = runner + ["--generate", excel, "--image-dir", image_dir, "--char-dir", P.char_dir,
                        "--model", a.emf_model, "--ratio", ratio or EMF_RATIOS[IMAGE_RATIO], "--batch", str(a.batch),
                        "--browser", a.browser,
                        "--project-store", os.path.join(P.dir, f"{P.stem}_flow_project.json")]
        if a.upscale:
            cmd += ["--upscale", a.upscale]
    if a.headless:
        cmd.append("--headless")
    if a.account:
        cmd += ["--account", a.account]
    if only:
        cmd += ["--only", ",".join(str(n) for n in sorted(only))]
    if dry_run:
        cmd.append("--dry-run")
    return cmd


def emf_login(P, log_prefix):
    script = FLOWUI_SCRIPT if P.args.image_backend == "flowui" else EMF_SCRIPT
    P.log(f"{log_prefix} Google Flow 로그인 확인 ({script}) → 창이 뜨면 로그인, 저장된 로그인이 있으면 자동 통과")
    login_cmd = flow_tool_runner(script) + ["--login"]
    if P.args.image_backend != "flowui":
        login_cmd += ["--browser", P.args.browser]
    if P.args.account:
        login_cmd += ["--account", P.args.account]
    rc, _ = run_streamed(login_cmd, cwd=TOOLS_DIR, log=P.log)
    if rc != 0:
        raise PipelineError("Google Flow 로그인 실패. `python Vrew_Auto_Pipeline_V1_0.py login` 으로 먼저 로그인하세요.")


def emf_generate(P, excel, image_dir, *, only=None, ratio=None, dry_run=False, log_prefix="[images]"):
    """이멀플 CLI 로 생성한다. 로그인이 없으면 로그인 창을 띄운 뒤 한 번 더, 한도 소진이면 중단."""
    if P.args.image_backend == "emf" and not dry_run and not P.args.headless:
        P.log(f"{log_prefix} 이멀플이 Chrome 창을 띄운 채 생성합니다 — 그 창을 닫거나 건드리지 마세요 (창 없이 돌리려면 --headless)")
    cmd = emf_command(P, excel, image_dir, only=only, ratio=ratio, dry_run=dry_run)
    rc, out = run_streamed(cmd, cwd=TOOLS_DIR, log=P.log)
    if dry_run:
        return out
    if any(k in out for k in EMF_LOGIN_NEEDED):
        emf_login(P, log_prefix)
        rc, out = run_streamed(emf_command(P, excel, image_dir, only=only, ratio=ratio), cwd=TOOLS_DIR, log=P.log)
    if "한도 소진" in out or "크레딧 부족" in out:
        raise PipelineError("Google Flow 계정 한도/크레딧 소진. 다른 계정으로 --account 를 지정하거나 크레딧 충전 후 다시 실행하세요.")
    if "연속 5회 실패" in out:
        P.log(f"{log_prefix} 생성기가 연속 실패로 멈췄습니다 — 남은 번호는 다음 회차에서 다시 시도합니다 (_flowui_debug/ 스크린샷 참고)")
    return out


# ── STAGE 3  characters — 캐릭터 참조 이미지(시트) 자동 생성 ──

REF_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}      # 이멀플 find_character_images 와 동일


def character_ref_map(folder):
    """{소문자 이름: 경로} — 이멀플과 같은 규칙(파일명 stem == 등장인물 이름, 대소문자 무시)"""
    m = {}
    if os.path.isdir(folder):
        for f in sorted(os.listdir(folder)):
            if os.path.splitext(f)[1].lower() in REF_IMAGE_EXTS:
                m[os.path.splitext(f)[0].lower()] = os.path.join(folder, f)
    return m


def storyboard_character_names(rows):
    seen = {}
    for r in rows:
        for n in (r["chars"] or "").split(","):
            n = n.strip()
            if n and n.lower() not in seen:
                seen[n.lower()] = n
    return list(seen.values())


def character_sheet_prompt(analysis, ch):
    costume = ""
    tl = ch.get("costume_timeline") or []
    if tl and isinstance(tl[0], dict):
        costume = str(tl[0].get("costume", "")).strip()
    parts = [str(analysis.get("style_prompt", "")).strip(), str(ch.get("description", "")).strip(), costume,
             "character reference sheet, single person, standing, full body, front view, looking at the camera, "
             "neutral expression, plain light gray studio background filling entire frame, soft even lighting, "
             "no text, no watermark"]
    return ", ".join(p.rstrip(",. ") for p in parts if p)


def stage_characters(P, rows):
    names = storyboard_character_names(rows)
    if not names:
        P.log("[characters] 등장인물이 없는 스토리보드 — 건너뜀")
        P.mark("characters", generated=0)
        return
    if P.args.image_backend == "flowui":
        P.log("[characters] FlowUI 백엔드는 참조 이미지 첨부를 아직 지원하지 않아 캐릭터 시트를 만들지 않습니다 "
              "(프롬프트의 캐릭터 묘사 블록으로 일관성 유지)")
        P.mark("characters", generated=0, skipped="flowui")
        return
    project_char = os.path.join(P.dir, "characters")
    common_char = os.path.join(SCRIPT_DIR, "characters")
    # 참조 이미지 후보: 공용 characters/ → (--char-dir 또는 현재 char_dir) → 프로젝트 characters/ (뒤가 우선).
    # 프로젝트 폴더가 이미 있어도 공용 폴더에 나중에 넣은 <이름>.png 를 놓치지 않도록 항상 공용 폴더를 본다.
    src_dirs = [d for d in dict.fromkeys([common_char, P.char_dir])
                if os.path.abspath(d) != os.path.abspath(project_char)]
    have = {}
    for d in src_dirs + [project_char]:
        have.update(character_ref_map(d))
    missing = [n for n in names if n.lower() not in have]
    P.log(f"[characters] 등장인물 {len(names)}명: {', '.join(names)} | 참조 이미지 있음 {len(names) - len(missing)}명, 없음 {len(missing)}명")

    def sync_project_refs():
        """공용/지정 폴더의 등장인물 이미지를 프로젝트 characters/ 로 복사하고 그 폴더를 char_dir 로 쓴다(images 단계는 한 폴더만 본다)."""
        os.makedirs(project_char, exist_ok=True)
        for d in src_dirs:
            refs = character_ref_map(d)
            for n in names:
                src = refs.get(n.lower())
                if src:
                    dst = os.path.join(project_char, os.path.basename(src))
                    if not os.path.exists(dst):
                        shutil.copy2(src, dst)
        P.char_dir = project_char

    if not missing:
        sync_project_refs()
        P.mark("characters", generated=0, names=names, char_dir=P.char_dir)
        return
    if not os.path.exists(P.analysis_json):
        sync_project_refs()
        P.log(f"[characters] 분석 JSON 이 없어(수동 스토리보드) 자동 생성을 건너뜁니다 → {P.char_dir} 에 <이름>.png 를 직접 넣으면 첨부됩니다: {', '.join(missing)}")
        P.mark("characters", generated=0, skipped=missing)
        return
    with open(P.analysis_json, "r", encoding="utf-8") as f:
        analysis = json.load(f)
    desc_by = {str(c.get("name", "")).lower(): c for c in analysis.get("characters", [])}
    todo = []
    for n in missing:
        ch = desc_by.get(n.lower())
        if not ch or not ch.get("description"):
            P.log(f"[characters]   {n}: 분석 JSON 에 외형 묘사가 없어 건너뜀 (참조 이미지 직접 추가 가능)")
            continue
        todo.append((n, character_sheet_prompt(analysis, ch)))
    if not todo:
        sync_project_refs()
        P.mark("characters", generated=0, skipped=missing)
        return

    # 생성 결과는 프로젝트 characters/ 에 둔다 (공용 characters/ 의 해당 인물 이미지는 복사해 함께 사용)
    sync_project_refs()

    P.log(f"[characters] 캐릭터 시트 {len(todo)}장 생성: {', '.join(n for n, _ in todo)} → {P.sheet_dir}")
    if P.args.dry_run:
        for n, prompt in todo:
            P.log(f"[characters]   [dry-run] {n}: {prompt[:120]}...")
        return
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["이미지 번호", "등장인물", "프롬프트"])          # 등장인물은 비워 둔다 (에셋 첨부 없이 생성)
    for i, (n, prompt) in enumerate(todo, 1):
        ws.append([i, "", prompt])
    wb.save(P.character_xlsx)
    os.makedirs(P.sheet_dir, exist_ok=True)
    # 이전 시트 파일이 남아 있으면 번호가 섞이므로 비운다
    for f in os.listdir(P.sheet_dir):
        if re.match(r"^\d+_", f):
            os.remove(os.path.join(P.sheet_dir, f))
    emf_generate(P, P.character_xlsx, P.sheet_dir, ratio=EMF_RATIOS[P.args.character_ratio], log_prefix="[characters]")
    made = []
    sheets = scan_media(P.sheet_dir)
    for i, (n, _) in enumerate(todo, 1):
        files = sheets.get(i, {}).get("image", [])
        if not files:
            P.log(f"[characters]   {n}: 생성 실패 (시트 없음)")
            continue
        src = files[-1]
        dst = os.path.join(project_char, n + os.path.splitext(src)[1].lower())
        shutil.copy2(src, dst)
        made.append(n)
        P.log(f"[characters]   {n} → {os.path.basename(dst)}")
    P.mark("characters", generated=len(made), names=made, char_dir=P.char_dir)
    if len(made) < len(todo):
        P.log(f"[characters] 일부 인물은 참조 이미지 없이 진행합니다: {[n for n, _ in todo if n not in made]}")


# ── STAGE 4  images ──

def stage_images(P, rows):
    expected = unique_image_numbers(rows)
    missing = missing_image_numbers(P, rows)
    P.log(f"[images] 필요한 이미지 {len(expected)}장, 이미 있음 {len(expected) - len(missing)}장, 생성 대상 {len(missing)}장 → {P.image_dir}")
    if not missing:
        P.mark("images", generated=0, missing=[])
        return
    os.makedirs(P.image_dir, exist_ok=True)
    if P.args.image_backend == "genspark":
        return run_genspark(P, rows)

    refs = character_ref_map(P.char_dir)
    P.log(f"[images] 캐릭터 참조 폴더: {P.char_dir} ({len(refs)}개)")
    for round_no in (1, 2):
        emf_generate(P, P.storyboard_xlsx, P.image_dir, only=set(missing), dry_run=P.args.dry_run)
        if P.args.dry_run:
            return
        missing = missing_image_numbers(P, rows)
        if not missing:
            break
        P.log(f"[images] {round_no}차 후 미생성 {len(missing)}장: {missing[:20]}{' ...' if len(missing) > 20 else ''}")
    P.mark("images", missing=missing)
    if missing:
        msg = f"[images] 끝내 생성되지 않은 이미지 {len(missing)}장: {missing}"
        if P.args.strict:
            raise PipelineError(msg + " (--strict)")
        P.log(msg + " → Vrew 연결 시 앞 번호 이미지로 대체(fallback)됩니다")


def run_genspark(P, rows):
    runner = find_tool_runner(GENSPARK_SCRIPT)
    if runner[0] == sys.executable and not python_has("selenium", "openpyxl"):
        raise PipelineError("Genspark 실행에 selenium 이 필요합니다: pip install selenium undetected-chromedriver openpyxl")
    cmd = runner + ["--cli", "--excel", P.storyboard_xlsx, "--char-dir", P.char_dir, "--browser", P.args.browser,
                    "--login-wait", str(P.args.login_wait)]
    if P.args.dry_run:
        cmd.append("--dry-run")
    rc, out = run_streamed(cmd, cwd=TOOLS_DIR, log=P.log)
    # Genspark 는 도구 폴더의 Image/ 에 저장한다 → 프로젝트 Image/ 로 동기화
    src = os.path.join(TOOLS_DIR, "Image")
    if os.path.abspath(src) != os.path.abspath(P.image_dir) and os.path.isdir(src):
        copied = 0
        for num, entry in scan_media(src).items():
            for f in entry["image"]:
                dst = os.path.join(P.image_dir, os.path.basename(f))
                if not os.path.exists(dst):
                    shutil.copy2(f, dst)
                    copied += 1
        P.log(f"[images] Genspark 결과 {copied}개 파일을 {P.image_dir} 로 복사")
    missing = missing_image_numbers(P, rows)
    P.mark("images", backend="genspark", missing=missing)
    if missing:
        P.log(f"[images] 미생성 이미지 {len(missing)}장: {missing}")


# ══════════════════════════════════════════════════════════════
#  STAGE 5  connect — Vrew 에 이미지(·직접 넣은 동영상) 연결
# ══════════════════════════════════════════════════════════════

def kenburns_rows(rows, mode):
    """connect 직전 켄번즈 열 조정. auto=스토리보드 값 그대로(Claude 지정) / random=이미지마다 무작위(마지막 이미지는 none) / none=효과 없음."""
    if mode not in ("random", "none"):
        return rows
    import random
    out = [dict(r) for r in rows]
    firsts = [r for r in out if str(r.get("prompt") or "").strip()]      # 이미지 그룹의 첫 행(켄번즈는 여기에만 적는다)
    for r in out:
        r["kenburns"] = ""
    if mode == "random":
        for i, r in enumerate(firsts):
            r["kenburns"] = "none" if i == len(firsts) - 1 else random.choice(KENBURNS_VALUES)
    return out


def stage_connect(P, rows):
    from Vrew_Connector_V1_8 import add_images_to_vrew
    media = scan_media(P.image_dir)
    expected = unique_image_numbers(rows)
    n_img = sum(1 for n in expected if n in media and media[n]["image"])
    n_vid = sum(1 for n in expected if n in media and media[n]["video"])
    if not media:
        raise PipelineError(f"{P.image_dir} 에 연결할 미디어가 없습니다.")
    P.log(f"[connect] 이미지 번호 {len(expected)}개 중 이미지 {n_img}개, 직접 넣은 동영상 {n_vid}개 준비됨 (같은 번호는 동영상 우선)")
    if P.args.dry_run:
        P.log("[connect] [dry-run] Vrew 파일을 쓰지 않습니다")
        return None
    out = P.args.output or os.path.join(P.dir, f"{P.stem}_이미지연결_{datetime.now().strftime('%Y%m%d_%H%M')}.vrew")
    kb_mode = getattr(P.args, "kenburns", "auto") or "auto"
    sb_xlsx = P.storyboard_xlsx
    if kb_mode in ("random", "none"):
        sb_xlsx = os.path.join(P.dir, f"{P.stem}_스토리보드_연결용.xlsx")   # 원본 스토리보드는 건드리지 않는다
        write_storyboard_xlsx(sb_xlsx, kenburns_rows(rows, kb_mode))
        P.log(f"[connect] 켄번즈 {'무작위' if kb_mode == 'random' else '없음'} → {os.path.basename(sb_xlsx)}")
    else:
        P.log("[connect] 켄번즈: 스토리보드 값 그대로 (Claude 지정)")
    add_images_to_vrew(P.vrew, sb_xlsx, P.image_dir, out,
                       crop_to_fill=not P.args.no_crop, use_fallback=not P.args.no_fallback,
                       log_callback=lambda m: P.log("  [Vrew] " + m))
    P.log(f"[connect] 완료 → {out}")
    P.mark("connect", output=out, images=n_img, videos=n_vid)
    return out


# ══════════════════════════════════════════════════════════════
#  명령
# ══════════════════════════════════════════════════════════════

PARSER = None      # main() 이 채운다 — 설정 파일 적용 때 '인자가 기본값인지' 판단용

# 스튜디오 UI 설정 파일 키 → argparse dest. (character_mode 는 값에 따라 플래그로 풀린다)
SETTINGS_DIRECT = ("style_prompt", "target_images", "block_seconds", "split_mode", "front_minutes", "front_per_min", "front_rest_per_min",
                   "emf_model", "batch", "account", "headless", "upscale", "character_ratio", "kenburns")


def apply_project_settings(args):
    """projects/<이름>/<이름>_settings.json(스튜디오 UI 저장값)을 args 에 적용한다. CLI 로 명시한 인자(기본값과 다른 값)가 우선.
    반환: 적용한 키 목록(로그용)."""
    path = os.path.join(os.path.dirname(os.path.abspath(args.vrew)),
                        f"{Path(args.vrew).stem}_settings.json")
    if not os.path.exists(path):
        return []
    try:
        with open(path, "r", encoding="utf-8") as f:
            s = json.load(f)
    except Exception as e:
        print(f"[settings] {os.path.basename(path)} 읽기 실패 — 무시: {e}")
        return []
    if not isinstance(s, dict):
        return []

    rp = getattr(PARSER, "run_parser", None)     # run 서브파서의 기본값과 비교해야 한다(상위 파서는 그 인자를 모른다)

    def is_default(dest):
        return rp is None or getattr(args, dest, None) == rp.get_default(dest)

    applied = []
    for k in SETTINGS_DIRECT:
        if k in s and hasattr(args, k) and is_default(k):
            v = s[k]
            setattr(args, k, v)
            applied.append(f"{k}={v!r}" if v != "" else f"{k}=''")
    mode = s.get("character_mode")
    if mode == "manual" and is_default("no_characters"):
        args.no_characters = True
        applied.append("no_characters(직접 준비)")
    return applied


def cmd_run(args):
    applied = apply_project_settings(args)
    P = Project(args.vrew, args)
    if applied:
        P.log(f"[settings] {os.path.basename(P.settings_json)} 적용: {', '.join(applied)}")
    if args.only:
        first = last = args.only
    else:
        first, last = args.from_stage, args.to_stage
    lo, hi = STAGES.index(first), STAGES.index(last)
    if lo > hi:
        raise PipelineError(f"--from {first} 가 --to {last} 보다 뒤입니다.")
    want_chars = not args.no_characters or args.only == "characters"
    plan = [s for i, s in enumerate(STAGES) if lo <= i <= hi
            and (s != "characters" or want_chars)]
    P.log(f"═══ 금광롱폼 Vrew 자동 파이프라인 v{__version__} ═══")
    P.log(f"입력: {P.vrew}")
    P.log(f"단계: {' → '.join(plan)}" + ("  [dry-run]" if args.dry_run else ""))
    os.makedirs(P.image_dir, exist_ok=True)

    rows = None
    try:
        return _run_stages(P, args, plan)
    finally:
        P.log("─── 실행 종료 ───")      # 스튜디오 UI 가 '진행 중' 판정을 끝내는 표식(부분 실행·게이트 정지에도 찍힌다)


def _run_stages(P, args, plan):
    rows = None
    for stage in plan:
        P.log(f"─── STEP {STAGES.index(stage)}/{len(STAGES) - 1}: {stage} ───")
        if stage == "login":
            stage_login(P, plan)
        elif stage == "extract":
            stage_extract(P)
        elif stage == "storyboard":
            rows = stage_storyboard(P)
        else:
            if rows is None:
                if not os.path.exists(P.storyboard_xlsx):
                    raise PipelineError(f"스토리보드가 없습니다: {P.storyboard_xlsx} (storyboard 단계를 먼저 실행)")
                rows, errors, _ = validate_storyboard(read_table_rows(P.storyboard_xlsx), P.clips(), strict_contiguous=False)
                if rows is None:
                    raise PipelineError("스토리보드 검증 실패: " + "; ".join(errors[:5]))
            if stage == "characters":
                stage_characters(P, rows)
            elif stage == "images":
                stage_images(P, rows)
            elif stage == "connect":
                out = stage_connect(P, rows)
                if out:
                    P.log("═══ 완료 ═══")
                    P.log(f"결과 Vrew: {out}")
    return 0


def cmd_status(args):
    P = Project(args.vrew, args)
    if getattr(args, "json", False):
        rows = read_table_rows(P.storyboard_xlsx) if os.path.exists(P.storyboard_xlsx) else []
        imgs = unique_image_numbers(rows) if rows else []
        media = scan_media(P.image_dir)
        print(json.dumps({
            "vrew": P.vrew, "dir": P.dir,
            "clips": len(P.clips()) if os.path.exists(P.clips_xlsx) else 0,
            "seconds": P.total_seconds() if os.path.exists(P.clips_xlsx) else 0,
            "storyboard_images": len(imgs),
            "images_have": sum(1 for n in imgs if n in media and media[n]["image"]),
            "videos_have": sum(1 for n in imgs if n in media and media[n]["video"]),
            "missing": [n for n in imgs if n not in media],
            "results": [str(p) for p in sorted(Path(P.dir).glob(f"{P.stem}_이미지연결_*.vrew"))],
            "stages": P.state.get("stages", {}),
        }, ensure_ascii=False, indent=2))
        return 0
    print(f"프로젝트: {P.vrew}")

    def line(label, path, extra=""):
        mark = "✅" if os.path.exists(path) else "⚪"
        print(f"  {mark} {label}: {os.path.basename(path)} {extra}")

    line("클립정보", P.clips_xlsx, f"({len(P.clips())}클립, {fmt_duration(P.total_seconds())})" if os.path.exists(P.clips_xlsx) else "")
    line("분석 JSON", P.analysis_json)
    extra = ""
    rows = None
    if os.path.exists(P.storyboard_xlsx):
        rows = read_table_rows(P.storyboard_xlsx)
        imgs = unique_image_numbers(rows)
        extra = f"(이미지 {len(imgs)}장)"
    line("스토리보드", P.storyboard_xlsx, extra)
    line("수동 JSON", P.storyboard_json)
    line("수동 CSV(구)", P.storyboard_csv)
    media = scan_media(P.image_dir)
    n_img = sum(1 for e in media.values() if e["image"])
    n_vid = sum(1 for e in media.values() if e["video"])
    print(f"  📁 Image/: 이미지 {n_img}개 번호, 동영상 {n_vid}개 번호")
    if rows:
        miss = [n for n in unique_image_numbers(rows) if n not in media]
        print(f"     미생성 이미지: {len(miss)}장 {miss[:30]}{' ...' if len(miss) > 30 else ''}")
    outs = sorted(Path(P.dir).glob(f"{P.stem}_이미지연결_*.vrew"))
    print(f"  🎬 결과 Vrew: {len(outs)}개" + (f" (최신: {outs[-1].name})" if outs else ""))
    for k, v in P.state.get("stages", {}).items():
        print(f"  · {k}: {v}")
    return 0


STAGE_KO = {"login": "로그인", "extract": "클립 추출", "storyboard": "스토리보드", "characters": "캐릭터 시트",
            "images": "이미지 생성", "connect": "Vrew 연결"}


def progress_line(P):
    """진행 한 줄 요약 — 클로드코드 세션이 30초마다 대화창에 옮겨 적는다(사용자 지시 2026-09-12).
    반환 (문장, 끝났는지). '끝'은 로그의 마지막 STEP 뒤에 '실행 종료' 표식이 있을 때, 또는 로그가 10분 넘게 멈췄을 때."""
    if not os.path.exists(P.log_file):
        return "아직 시작 전 (로그 없음)", False
    with open(P.log_file, "r", encoding="utf-8", errors="replace") as f:
        lines = [l.rstrip("\n") for l in f.readlines()[-400:]]
    last_step = max((i for i, l in enumerate(lines) if "─── STEP" in l), default=-1)
    last_end = max((i for i, l in enumerate(lines) if "실행 종료" in l), default=-1)
    finished = last_end > last_step
    stage = ""
    if last_step >= 0:
        m = re.search(r"STEP \d+/\d+: (\w+)", lines[last_step])
        stage = m.group(1) if m else ""
    tail = [l for l in lines if l.strip() and "───" not in l]
    last = re.sub(r"^\[[^\]]+\] ", "", tail[-1]) if tail else ""
    detail = last[:90]
    try:
        if stage == "characters" and os.path.exists(P.analysis_json):
            with open(P.analysis_json, "r", encoding="utf-8") as f:
                names = [c.get("name", "") for c in (json.load(f).get("characters") or [])]
            have = sum(1 for n in names if any(os.path.exists(os.path.join(d, n + ext))
                                                 for d in (P.char_dir, os.path.join(SCRIPT_DIR, "characters"))
                                                 for ext in (".png", ".jpg", ".jpeg", ".webp")))
            detail = f"캐릭터 시트 {have}/{len(names)}명 · {last[:60]}"
        elif stage in ("images", "connect") and os.path.exists(P.storyboard_xlsx):
            imgs = unique_image_numbers(read_table_rows(P.storyboard_xlsx))
            media = scan_media(P.image_dir)
            have = sum(1 for n in imgs if n in media)
            detail = f"이미지 {have}/{len(imgs)}장 · {last[:60]}"
        elif stage == "extract" and os.path.exists(P.clips_xlsx):
            detail = f"클립 {len(P.clips())}개, {fmt_duration(P.total_seconds())}"
    except Exception as e:  # 요약은 실패해도 실행에 영향 없어야 한다
        detail = f"{last[:70]} (요약 실패: {type(e).__name__})"
    stale = time.time() - os.path.getmtime(P.log_file)
    if finished:
        state = "실행 종료"
    elif stale >= 600:
        state = f"멈춤? 로그 {int(stale // 60)}분째 갱신 없음"
    else:
        state = "진행 중"
    return f"{STAGE_KO.get(stage, stage or '-')} · {state} · {detail}", finished or stale >= 600


def cmd_progress(args):
    """한 줄 진행 요약. --watch N 이면 N초마다 반복 출력하고, 실행이 끝나면(실행 종료 표식) 스스로 끝난다.
    시작 직후라 아직 새 실행이 로그에 안 찍혔으면 최대 3분 기다린다."""
    P = Project(args.vrew, args)
    every = float(args.watch or 0)
    armed = False
    waited = 0.0
    while True:
        line, done = progress_line(P)
        if every and not armed:
            if done:                                   # 옛 실행의 종료 표식만 보인다 → 새 실행이 시작되길 기다림
                if waited >= 180:
                    print(f"[{time.strftime('%H:%M:%S')}] 새 실행이 시작되지 않았습니다 — 감시 종료", flush=True)
                    return 0
                time.sleep(min(every, 5.0))
                waited += min(every, 5.0)
                continue
            armed = True
        print(f"[{time.strftime('%H:%M:%S')}] {line}", flush=True)
        if not every or done:
            return 0
        time.sleep(every)


def cmd_ui(args):
    """금광 스튜디오(tools/Studio_UI_V1_0.py)를 띄운다 — 이미 떠 있으면 브라우저 탭만 연다(멱등). 서버는 분리 프로세스로 남는다."""
    import socket
    import webbrowser
    script = os.path.join(TOOLS_DIR, "Studio_UI_V1_0.py")
    if not os.path.isfile(script):
        raise PipelineError(f"스튜디오 UI 스크립트가 없습니다: {script}")
    port = int(args.port)
    url = f"http://127.0.0.1:{port}/"

    def up():
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.settimeout(0.4)
        try:
            return s.connect_ex(("127.0.0.1", port)) == 0
        finally:
            s.close()

    if not up():
        env = dict(os.environ, VREW_PIPELINE_NO_REEXEC="1", PYTHONIOENCODING="utf-8", PYTHONUTF8="1")
        flags = (subprocess.DETACHED_PROCESS | subprocess.CREATE_NEW_PROCESS_GROUP) if os.name == "nt" else 0
        logf = open(os.path.join(SCRIPT_DIR, "studio_ui.log"), "ab")
        subprocess.Popen([sys.executable, script, "--no-open", "--port", str(port)], cwd=SCRIPT_DIR, env=env,
                         stdout=logf, stderr=logf, stdin=subprocess.DEVNULL, creationflags=flags, close_fds=True)
        for _ in range(30):
            if up():
                break
            time.sleep(0.2)
        print(f"스튜디오 UI 시작: {url}" if up() else f"스튜디오 UI 가 뜨지 않았습니다 — studio_ui.log 확인")
    else:
        print(f"스튜디오 UI 실행 중: {url}")
    if not args.no_open:
        webbrowser.open(url)
    return 0 if up() else 1


def cmd_login(args):
    script = FLOWUI_SCRIPT if getattr(args, "backend", "emf") == "flowui" else EMF_SCRIPT
    cmd = flow_tool_runner(script) + ["--login"]
    if args.account:
        cmd += ["--account", args.account]
    rc, _ = run_streamed(cmd, cwd=TOOLS_DIR)
    return rc


def build_parser():
    p = argparse.ArgumentParser(description="금광롱폼 Vrew 자동 파이프라인 — TTS Vrew 파일 → 이미지 연결 Vrew 파일",
                                formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__.split("사용법", 1)[1])
    sub = p.add_subparsers(dest="cmd")

    i = sub.add_parser("init", help="projects/<날짜>_<브루파일명>/ 프로젝트 폴더를 만들고 .vrew 를 넣는다")
    i.add_argument("vrew", help="원본 .vrew 파일")
    i.add_argument("--date", default="", help="폴더 접두 날짜 (기본: 오늘 YYYYMMDD)")
    i.add_argument("--name", default="", help="폴더 이름 직접 지정 (기본: <날짜>_<브루파일명>)")
    i.add_argument("--move", action="store_true", help="복사 대신 이동")

    r = sub.add_parser("run", help="파이프라인 실행")
    r.add_argument("vrew", help="TTS 음성이 들어간 .vrew 파일")
    r.add_argument("--from", dest="from_stage", default="login", choices=STAGES)
    r.add_argument("--no-login-check", action="store_true", help="시작 시 로그인 확인 단계 생략")
    r.add_argument("--to", dest="to_stage", default="connect", choices=STAGES)
    r.add_argument("--only", choices=STAGES, help="한 단계만 실행")
    r.add_argument("--force", action="store_true", help="기존 클립정보/스토리보드/분석을 무시하고 다시 만든다")
    r.add_argument("--dry-run", action="store_true", help="Claude/브라우저/Vrew 쓰기 없이 계획만 확인")
    r.add_argument("--strict", action="store_true", help="이미지가 하나라도 안 만들어지면 중단")
    # storyboard
    r.add_argument("--storyboard-backend", default="auto", choices=["auto", "api", "manual", "claude-cli"],
                   metavar="{auto,api,manual}",
                   help="auto=API 키 있으면 api, 없으면 manual(기본) / api=Anthropic API / manual=요청 txt 만들고 종료(코드 3)")
    r.add_argument("--model", default="", help="Claude 모델 (API: claude-opus-5 …)")
    r.add_argument("--target-images", type=int, default=0, help="목표 이미지 수 (기본: 총 길이 ÷ --block-seconds)")
    r.add_argument("--block-seconds", type=float, default=DEFAULT_BLOCK_SECONDS, help="이미지 한 장당 목표 길이(초)")
    r.add_argument("--chunk-size", type=int, default=DEFAULT_CHUNK_SIZE, help="Claude 1회 호출당 클립 수")
    r.add_argument("--style-prompt", default="", help="스타일 고정용 프롬프트 (없으면 스타일 표현 금지) — 마법사 단계 2")
    r.add_argument("--split-mode", default="smart", choices=["smart", "uniform", "front"],
                   help="이미지 분할: smart=장면 전환 우선(기본) / uniform=--block-seconds 간격 균등 분할 / front=초반 집중 — 마법사 단계 3")
    r.add_argument("--front-minutes", type=float, default=10.0, help="front: 초반 집중 구간 길이(분) — 라움튜브 front:M:A,rest:B 의 M")
    r.add_argument("--front-per-min", type=float, default=3.0, help="front: 초반 구간 분당 장수 (A)")
    r.add_argument("--front-rest-per-min", type=float, default=1.0, help="front: 이후 구간 분당 장수 (B)")
    # connect
    r.add_argument("--kenburns", default="auto", choices=["auto", "random", "none"],
                   help="Vrew 연결 때 켄번즈(줌·팬) 효과: auto=스토리보드 값(Claude 지정, 기본) / random=이미지마다 무작위 / none=없음")
    # characters
    r.add_argument("--no-characters", action="store_true", help="캐릭터 참조 이미지 자동 생성 단계 생략")
    r.add_argument("--character-ratio", default="1:1", choices=list(EMF_RATIOS), help="캐릭터 시트 비율")
    # images
    r.add_argument("--image-backend", default="emf", choices=["emf", "flowui", "genspark"],
                   help="emf=이멀플 v5.6 flow.google.com RPC(기본, 배치·참조 첨부) / flowui=화면 자동화 폴백 / genspark")
    r.add_argument("--emf-model", default="Nano Banana 2", help="모델명 (Nano Banana 2 / Nano Banana 2 Lite / Nano Banana Pro)")
    r.add_argument("--upscale", default="", choices=["", "2K", "4K"])
    r.add_argument("--batch", type=int, default=3, help="이멀플 한 호출당 장수")
    r.add_argument("--headless", action="store_true")
    r.add_argument("--browser", default="Chrome", help="이멀플/Genspark/브라우저 종류 (Chrome/Edge/…)")
    r.add_argument("--account", default="", help="이멀플 계정 슬롯 id")
    r.add_argument("--char-dir", default="", help="캐릭터 참조 이미지 폴더 (기본: 프로젝트 또는 스크립트 폴더의 characters/)")
    r.add_argument("--login-wait", type=int, default=300, help="브라우저 로그인 대기(초)")
    # connect
    r.add_argument("--output", default="", help="결과 .vrew 경로 (기본: <이름>_이미지연결_<시각>.vrew)")
    r.add_argument("--no-crop", action="store_true", help="이미지를 16:9 로 자르지 않음")
    r.add_argument("--no-fallback", action="store_true", help="없는 번호를 앞 이미지로 대체하지 않음")

    s = sub.add_parser("status", help="진행 상태")
    s.add_argument("vrew")
    s.add_argument("--json", action="store_true", help="JSON 으로 출력 (스튜디오 UI·스크립트용)")
    g = sub.add_parser("progress", help="진행 한 줄 요약 (--watch N: N초마다 반복, 실행 종료 표식이 찍히면 끝)")
    g.add_argument("vrew")
    g.add_argument("--watch", type=float, default=0, help="N초마다 반복 출력 (클로드코드가 대화창에 옮겨 적는 용도)")

    u = sub.add_parser("ui", help="금광 스튜디오 UI 띄우기 (이미 떠 있으면 브라우저 탭만)")
    u.add_argument("--port", type=int, default=7788)
    u.add_argument("--no-open", action="store_true")
    p.run_parser = r      # apply_project_settings 가 run 인자의 기본값을 조회하는 데 쓴다

    l = sub.add_parser("login", help="Google Flow 로그인 (프로필 저장)")
    l.add_argument("--account", default="")
    l.add_argument("--backend", default="emf", choices=["emf", "flowui"], help="로그인 창을 띄울 도구 (프로필은 공유됨)")
    return p


def cmd_init(args):
    """projects/<날짜>_<브루파일명>/ 폴더를 만들고 .vrew 를 그 안에 복사(또는 이동)한 뒤 새 경로를 출력한다."""
    src = os.path.abspath(args.vrew)
    if not os.path.isfile(src) or not src.lower().endswith(".vrew"):
        raise PipelineError(f".vrew 파일이 아닙니다: {src}")
    stem = Path(src).stem
    parent = os.path.dirname(src)
    if os.path.normcase(os.path.dirname(parent)) == os.path.normcase(PROJECTS_DIR):
        print(f"이미 projects/ 안의 프로젝트 폴더에 있습니다: {parent}")
        print(f"VREW={src}")
        return 0
    date = args.date or datetime.now().strftime("%Y%m%d")
    folder_name = args.name or f"{date}_{stem}"
    folder_name = re.sub(r'[\\/:*?"<>|]+', "_", folder_name).strip()
    dest_dir = os.path.join(PROJECTS_DIR, folder_name)
    dest = os.path.join(dest_dir, os.path.basename(src))
    os.makedirs(dest_dir, exist_ok=True)
    if os.path.exists(dest):
        try:
            same = os.path.samefile(src, dest)
        except OSError:
            same = False
        if not same and os.path.getsize(src) != os.path.getsize(dest):
            raise PipelineError(f"이미 다른 파일이 있습니다: {dest} (--name 으로 다른 폴더명을 지정하세요)")
        print(f"프로젝트 폴더에 이미 같은 파일이 있어 재사용합니다: {dest}")
    else:
        if args.move:
            shutil.move(src, dest)
            print(f"이동: {src} → {dest}")
        else:
            shutil.copy2(src, dest)
            print(f"복사: {src} → {dest}")
    print(f"프로젝트 폴더: {dest_dir}")
    print(f"VREW={dest}")
    return 0


def main(argv=None):
    reexec_in_runtime()
    apply_runtime_env()
    parser = build_parser()
    global PARSER
    PARSER = parser
    args = parser.parse_args(argv)
    if not args.cmd:
        parser.print_help()
        return 0
    try:
        if args.cmd == "ui":
            return cmd_ui(args)
        if args.cmd == "init":
            return cmd_init(args)
        if args.cmd == "run":
            return cmd_run(args)
        if args.cmd == "status":
            return cmd_status(args)
        if args.cmd == "progress":
            return cmd_progress(args)
        if args.cmd == "login":
            return cmd_login(args)
    except PipelineError as e:
        print(f"\n❌ {e}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\n사용자 중단", file=sys.stderr)
        return 130
    except Exception:
        traceback.print_exc()
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
